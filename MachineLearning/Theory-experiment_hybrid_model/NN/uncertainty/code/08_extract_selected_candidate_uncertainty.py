from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from uncertainty_common import load_config, save_formatted_workbook


RUN_ID, _CONFIG = load_config()
BASE_DIR = Path(__file__).resolve().parent.parent
RESULT_DIR = BASE_DIR / "result"

CANDIDATE_PATH = RESULT_DIR / "high_potential_candidate_recommendations.xlsx"
UNCERTAINTY_PATH = RESULT_DIR / "uncertainty_table.xlsx"
OUTPUT_PATH = RESULT_DIR / "selected_high_potential_candidates_with_formal_PI95.xlsx"

COMPOSITION_COLUMNS = ["Mn", "Fe", "Co", "Ni", "Zn"]
FORMAL_COLUMNS = [
    "predicted_potential",
    "sigma_ML_epi",
    "sigma_descriptor",
    "sigma_emp_residual",
    "sigma_total",
    "H95",
    "PI95_lower",
    "PI95_upper",
    "Zone",
]
EXPECTED_CANDIDATE_COUNT = 40
EXPECTED_STAGE_COUNTS = {"global_top10": 10, "zone_top": 7, "zone_diversity": 23}
EXPECTED_GUIDED_ZONE_COUNTS = {"Z1": 6, "Z2": 10, "Z3": 6, "Z4": 8}
CANDIDATE_BENCHMARK_E10_V = float(_CONFIG["candidate_benchmark_E10_V"])
AT_PERCENT_TOLERANCE = 1e-6


def require_columns(frame, columns, source_name):
    missing = [column for column in columns if column not in frame.columns]
    if missing:
        raise ValueError(f"{source_name} is missing columns: {missing}")


def composition_key(values):
    values = np.asarray(values, dtype=float)
    if not np.isfinite(values).all():
        raise ValueError(f"Composition contains non-finite values: {values.tolist()}")
    if not np.isclose(values.sum(), 1.0, atol=1e-8):
        raise ValueError(f"Composition does not sum to 1: {values.tolist()}")
    at_percent = values * 100.0
    rounded = np.rint(at_percent).astype(int)
    if np.max(np.abs(at_percent - rounded)) > AT_PERCENT_TOLERANCE:
        raise ValueError(f"Composition is not on the 1 at.% grid: {values.tolist()}")
    return tuple(int(value) for value in rounded)


def add_composition_keys(frame, source_name):
    out = frame.copy()
    out["_composition_key"] = [
        composition_key([row[column] for column in COMPOSITION_COLUMNS])
        for _, row in out.iterrows()
    ]
    if out["_composition_key"].duplicated().any():
        duplicates = out.loc[out["_composition_key"].duplicated(False), COMPOSITION_COLUMNS]
        raise ValueError(f"{source_name} contains duplicate compositions:\n{duplicates}")
    return out


def read_selected_candidates():
    if not CANDIDATE_PATH.exists():
        raise FileNotFoundError(
            f"Missing Step 07 candidate file: {CANDIDATE_PATH}. Run Step 07 after Steps 01-05."
        )
    candidates = pd.read_excel(CANDIDATE_PATH, sheet_name="Sheet1")
    require_columns(
        candidates,
        [
            *COMPOSITION_COLUMNS,
            "predicted_potential",
            "Reason",
            "selection_stage",
            "Region",
            "global_prediction_rank",
            "zone_rank",
            "global_rank",
            "output_order",
            "PI95_lower",
            "candidate_benchmark_E10_V",
            "candidate_eligible",
            "uncertainty_source",
        ],
        CANDIDATE_PATH.name,
    )
    if len(candidates) != EXPECTED_CANDIDATE_COUNT:
        raise ValueError(
            f"Expected {EXPECTED_CANDIDATE_COUNT} Step 07 candidates, found {len(candidates)}."
        )
    stage_counts = {
        str(key): int(value)
        for key, value in candidates["selection_stage"].value_counts().to_dict().items()
    }
    if stage_counts != EXPECTED_STAGE_COUNTS:
        raise ValueError(f"Step 07 stage counts are {stage_counts}; expected {EXPECTED_STAGE_COUNTS}.")

    global_top = candidates.loc[candidates["selection_stage"] == "global_top10"]
    if sorted(global_top["global_prediction_rank"].astype(int).tolist()) != list(range(1, 11)):
        raise ValueError("Global Top-10 must have global_prediction_rank values 1-10.")
    guided = candidates.loc[candidates["selection_stage"] != "global_top10"]
    guided_zone_counts = {
        str(key): int(value) for key, value in guided["Region"].value_counts().to_dict().items()
    }
    if guided_zone_counts != EXPECTED_GUIDED_ZONE_COUNTS:
        raise ValueError(
            f"The 30 uncertainty-guided zone counts are {guided_zone_counts}; "
            f"expected {EXPECTED_GUIDED_ZONE_COUNTS}."
        )
    if not (guided["PI95_lower"] < CANDIDATE_BENCHMARK_E10_V).all():
        raise ValueError(
            "Every uncertainty-guided candidate must satisfy "
            f"PI95_lower < {CANDIDATE_BENCHMARK_E10_V:.4f} V."
        )
    candidates = add_composition_keys(candidates, CANDIDATE_PATH.name)
    return candidates.sort_values("output_order", kind="mergesort").reset_index(drop=True)


def read_formal_uncertainty(candidate_keys):
    if not UNCERTAINTY_PATH.exists():
        raise FileNotFoundError(f"Missing formal uncertainty table: {UNCERTAINTY_PATH}")
    workbook = load_workbook(UNCERTAINTY_PATH, read_only=True, data_only=True)
    try:
        if "Sheet1" not in workbook.sheetnames:
            raise ValueError(f"{UNCERTAINTY_PATH.name} does not contain Sheet1.")
        worksheet = workbook["Sheet1"]
        headers = list(next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True)))
        required = [*COMPOSITION_COLUMNS, *FORMAL_COLUMNS]
        missing = [column for column in required if column not in headers]
        if missing:
            raise ValueError(f"{UNCERTAINTY_PATH.name}:Sheet1 is missing columns: {missing}")
        indices = {column: headers.index(column) for column in required}
        max_column = max(indices.values()) + 1
        matches = {}
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, max_col=max_column, values_only=True), start=2
        ):
            composition = [values[indices[column]] for column in COMPOSITION_COLUMNS]
            if any(value is None for value in composition):
                continue
            key = composition_key(composition)
            if key not in candidate_keys:
                continue
            if key in matches:
                raise ValueError(f"Composition {key} occurs more than once in {UNCERTAINTY_PATH.name}.")
            matches[key] = {
                "_composition_key": key,
                **{column: values[indices[column]] for column in COMPOSITION_COLUMNS},
                **{column: values[indices[column]] for column in FORMAL_COLUMNS},
                "uncertainty_table_row": row_number,
            }
            if len(matches) == len(candidate_keys):
                break
    finally:
        workbook.close()

    missing_keys = sorted(candidate_keys - set(matches))
    if missing_keys:
        raise ValueError(
            f"The following Step 07 candidates were not found in {UNCERTAINTY_PATH.name}: {missing_keys}"
        )
    return pd.DataFrame(matches.values())


def validate_formal_intervals(table):
    numeric_columns = FORMAL_COLUMNS[:-1]
    if table[numeric_columns].isna().any().any():
        raise ValueError("Formal uncertainty values contain missing data.")
    if not np.allclose(
        table["PI95_lower"], table["predicted_potential"] - table["H95"], atol=1e-12
    ):
        raise ValueError("PI95_lower does not equal predicted_potential - H95.")
    if not np.allclose(
        table["PI95_upper"], table["predicted_potential"] + table["H95"], atol=1e-12
    ):
        raise ValueError("PI95_upper does not equal predicted_potential + H95.")
    expected_total = np.sqrt(
        table["sigma_ML_epi"] ** 2
        + table["sigma_descriptor"] ** 2
        + table["sigma_emp_residual"] ** 2
    )
    if not np.allclose(table["sigma_total"], expected_total, atol=1e-12):
        raise ValueError("sigma_total does not equal the quadrature sum of its three components.")
    nonnegative = ["sigma_ML_epi", "sigma_descriptor", "sigma_emp_residual", "sigma_total", "H95"]
    if (table[nonnegative] < 0).any().any():
        raise ValueError("Formal uncertainty values must be non-negative.")


def selection_strategy(value):
    labels = {
        "global_top10": "Global predicted-potential Top-10",
        "zone_top": "Zone top",
        "zone_diversity": "ILR diversity",
    }
    return labels.get(str(value), str(value))


def build_output_tables(candidates, formal):
    candidate_columns = [
        "_composition_key",
        "Fe",
        "Co",
        "Ni",
        "Mn",
        "Zn",
        "Reason",
        "selection_stage",
        "Region",
        "global_prediction_rank",
        "zone_rank",
        "global_rank",
        "output_order",
        "candidate_benchmark_E10_V",
        "candidate_eligible",
        "formal_PI95_eligible",
        "ilr_distance_score",
        "optimistic_PI_score",
        "selection_score",
        "distance_reference",
        "uncertainty_source",
    ]
    candidate_columns = [column for column in candidate_columns if column in candidates.columns]
    candidate_metadata = candidates[candidate_columns].copy()
    formal_values = formal[["_composition_key", *FORMAL_COLUMNS, "uncertainty_table_row"]]
    merged = candidate_metadata.merge(
        formal_values, on="_composition_key", how="left", validate="one_to_one"
    ).sort_values("output_order", kind="mergesort")
    validate_formal_intervals(merged)

    reviewer_table = pd.DataFrame(
        {
            "Selection channel": merged["selection_stage"].map(selection_strategy),
            "Zone": merged["Zone"],
            "Fe": merged["Fe"],
            "Co": merged["Co"],
            "Ni": merged["Ni"],
            "Mn": merged["Mn"],
            "Zn": merged["Zn"],
            "Predicted potential (V)": merged["predicted_potential"],
            "95% empirical PI lower bound for E10 (V)": merged["PI95_lower"],
            "Experimentally validated potential (V)": np.nan,
        }
    )
    detail_columns = [
        "output_order",
        "Fe",
        "Co",
        "Ni",
        "Mn",
        "Zn",
        "Reason",
        "selection_stage",
        "Region",
        "global_prediction_rank",
        "zone_rank",
        "global_rank",
        "candidate_benchmark_E10_V",
        "candidate_eligible",
        "formal_PI95_eligible",
        "ilr_distance_score",
        "optimistic_PI_score",
        "selection_score",
        "distance_reference",
        "uncertainty_source",
        "predicted_potential",
        "sigma_ML_epi",
        "sigma_descriptor",
        "sigma_emp_residual",
        "sigma_total",
        "H95",
        "PI95_lower",
        "PI95_upper",
        "Zone",
        "uncertainty_table_row",
    ]
    uncertainty_detail = merged[detail_columns].reset_index(drop=True)
    metadata = pd.DataFrame(
        [
            {"key": "RUN_ID", "value": RUN_ID},
            {"key": "candidate_source", "value": f"result/{CANDIDATE_PATH.name}"},
            {"key": "formal_uncertainty_source", "value": f"result/{UNCERTAINTY_PATH.name}"},
            {"key": "formal_uncertainty_sheet", "value": "Sheet1"},
            {"key": "candidate_count", "value": EXPECTED_CANDIDATE_COUNT},
            {"key": "global_prediction_top_count", "value": 10},
            {"key": "uncertainty_guided_count", "value": 30},
            {"key": "H_source", "value": "formal uncertainty_table H95"},
            {"key": "composition_match", "value": "integer 1 at.% key in Mn/Fe/Co/Ni/Zn order"},
            {"key": "candidate_benchmark_E10_V", "value": CANDIDATE_BENCHMARK_E10_V},
            {
                "key": "candidate_selection",
                "value": "global predicted-potential Top-10 plus 30 formal-PI95/zone/ILR candidates",
            },
        ]
    )
    return reviewer_table, uncertainty_detail, metadata


def main():
    candidates = read_selected_candidates()
    candidate_keys = set(candidates["_composition_key"])
    formal = read_formal_uncertainty(candidate_keys)
    reviewer_table, uncertainty_detail, metadata = build_output_tables(candidates, formal)
    save_formatted_workbook(
        OUTPUT_PATH,
        {
            "reviewer_table": reviewer_table,
            "uncertainty_detail": uncertainty_detail,
            "metadata": metadata,
        },
    )
    print(f"Matched {len(uncertainty_detail)} Step 07 candidates.")
    print("H_source: formal uncertainty_table H95")
    print(f"Saved formal selected-candidate uncertainty to {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
