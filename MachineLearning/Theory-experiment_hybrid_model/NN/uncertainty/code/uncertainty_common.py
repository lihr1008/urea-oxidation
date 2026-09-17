import importlib.util
import json
import math
import pickle
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import torch
from sklearn.decomposition import PCA
from sklearn.model_selection import KFold
from sklearn.preprocessing import StandardScaler


ML_ROOT = Path(__file__).resolve().parents[4]
REPO_ROOT = ML_ROOT.parent
if str(REPO_ROOT) not in sys.path:
    sys.path.append(str(REPO_ROOT))

from nn_common import (  # noqa: E402
    DESCRIPTOR_COLUMNS,
    DescriptorPreModel,
    DescriptorToExperimentModel,
    RATIO_COLUMNS,
    load_xy,
    make_loader,
    metrics,
    read_choose_proportion,
    save_pickle,
    split_scale,
    train_model,
)


AD95_LABEL_INSIDE = "inside AD95"
AD95_LABEL_OUTSIDE = "outside AD95"
REQUIRED_REPEATABILITY_COLUMNS = RATIO_COLUMNS + ["E10_rep1", "E10_rep2", "E10_rep3"]
DFT_DESCRIPTOR_PROPERTY_MAP = {
    "\u0394E*OC(NH2)2": "ade_urea",
    "\u0394E*CO": "ade_CO",
    "\u0394q*OC(NH2)2": "e_urea",
    "\u0394q*CO": "e_CO",
}
SCALER_DESCRIPTOR_COLUMNS = ["\u0394E_urea", "\u0394E_CO", "\u0394e_urea", "\u0394q_CO"]
DESCRIPTOR_DISPLAY_NAMES = {
    "ade_urea": "\u0394E*OC(NH2)2",
    "ade_CO": "\u0394E*CO",
    "e_urea": "\u0394q*OC(NH2)2",
    "e_CO": "\u0394q*CO",
}
DESCRIPTOR_UNITS = {"ade_urea": "eV", "ade_CO": "eV", "e_urea": "e", "e_CO": "e"}
DESCRIPTOR_UNCERTAINTY_SOURCES = [
    "temperature",
    "force_field",
    "functional",
    "slab_size",
    "second_shell",
]
DESCRIPTOR_SOURCE_REPORT_ORDER = [
    "functional",
    "slab_size",
    "second_shell",
    "temperature",
    "force_field",
]
DESCRIPTOR_SOURCE_DISPLAY_NAMES = {
    "functional": "DFT functional",
    "slab_size": "Slab size",
    "second_shell": "Second coordination shell",
    "temperature": "MD temperature/time",
    "force_field": "Force field",
}
MD_RATIO_COLUMN = "Metal ratio (Fe:Co:Ni:Mn:Zn)"
MD_TEMPERATURE_COLUMNS = ["2000 K/1 ns", "1500 K/2 ns", "1000 K/3 ns"]


def load_config():
    config_path = Path(__file__).resolve().parent / "00_config.py"
    spec = importlib.util.spec_from_file_location("uncertainty_pipeline_config", config_path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module.RUN_ID, module.CONFIG


def to_builtin(value):
    if isinstance(value, dict):
        return {str(key): to_builtin(item) for key, item in value.items()}
    if isinstance(value, (list, tuple)):
        return [to_builtin(item) for item in value]
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, Path):
        return str(value)
    return value


def save_json(data, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(to_builtin(data), f, ensure_ascii=False, indent=2)


def load_pickle(path):
    with open(path, "rb") as f:
        return pickle.load(f)


def resolve_repo_path(path):
    path = Path(path)
    return path if path.is_absolute() else REPO_ROOT / path


def run_dirs(config):
    base_dir = Path(__file__).resolve().parent.parent
    model_dir = base_dir / "model_load" / config["run_id"]
    result_dir = base_dir / "result"
    result_dir.mkdir(parents=True, exist_ok=True)
    return model_dir, result_dir


def set_all_seeds(seed):
    np.random.seed(int(seed))
    torch.manual_seed(int(seed))
    if torch.cuda.is_available():
        torch.cuda.manual_seed_all(int(seed))


def _predict_scaled(model, x_scaled):
    model.eval()
    with torch.no_grad():
        return model(torch.tensor(x_scaled).to(torch.float32)).detach().numpy()


def _fit_premodel(seed, pre_config):
    set_all_seeds(seed)
    x_pre, y_pre = load_xy("metals_activity.xlsx", "metals", "activity")
    pre_data = split_scale(x_pre, y_pre, seed, pre_config["test_size"])
    premodel = DescriptorPreModel(
        pre_config["input_dim"], pre_config["hidden_units"], pre_config["output_dim"]
    )
    pre_history = train_model(
        premodel,
        make_loader(pre_data["x_train_scaled"], pre_data["y_train_scaled"], pre_config["batch_size"]),
        make_loader(pre_data["x_test_scaled"], pre_data["y_test_scaled"], pre_config["batch_size"]),
        pre_config,
    )
    train_pred = _predict_scaled(premodel, pre_data["x_train_scaled"])
    test_pred = _predict_scaled(premodel, pre_data["x_test_scaled"])
    pre_metric = metrics(
        pre_data["y_train"],
        pre_data["y_test"],
        pre_data["norm_y"].inverse_transform(train_pred),
        pre_data["norm_y"].inverse_transform(test_pred),
    )
    return premodel, pre_data, pre_history, pre_metric


def _fit_descriptor_to_z(
    premodel, pretrain_norm_x, x_train, y_train, x_eval, y_eval, train_config
):
    norm_y = StandardScaler().fit(y_train)
    x_train_scaled = pretrain_norm_x.transform(x_train)
    x_eval_scaled = pretrain_norm_x.transform(x_eval)
    y_train_scaled = norm_y.transform(y_train)
    y_eval_scaled = norm_y.transform(y_eval)
    with torch.no_grad():
        train_descriptor = premodel(torch.tensor(x_train_scaled).to(torch.float32)).detach()
        eval_descriptor = premodel(torch.tensor(x_eval_scaled).to(torch.float32)).detach()

    model = DescriptorToExperimentModel(
        train_config["input_dim"],
        train_config["hidden_units"],
        train_config["output_dim"],
        train_config["dropout"],
    )
    history = train_model(
        model,
        make_loader(train_descriptor, y_train_scaled, train_config["batch_size"]),
        make_loader(eval_descriptor, y_eval_scaled, train_config["batch_size"]),
        train_config,
    )
    train_pred = norm_y.inverse_transform(_predict_scaled(model, train_descriptor))
    eval_pred = norm_y.inverse_transform(_predict_scaled(model, eval_descriptor))
    metric = metrics(y_train, y_eval, train_pred, eval_pred)
    return {
        "model": model,
        "norm_y": norm_y,
        "history": history,
        "metric": metric,
        "eval_pred": eval_pred,
    }


def train_cv_seed_ensemble(config):
    model_dir, result_dir = run_dirs(config)
    model_dir.mkdir(parents=True, exist_ok=True)
    cv_dir = model_dir / "cv_seed_ensemble"
    cv_dir.mkdir(parents=True, exist_ok=True)
    x_exp, y_exp = load_xy("experiment.xlsx", "metals", "potential")
    grid = pd.DataFrame(read_choose_proportion(), columns=RATIO_COLUMNS)
    n_repeats = int(config.get("cv_n_repeats", 1))
    n_splits = int(config["cv_n_splits"])
    seed_per_fold = int(config["seed_per_fold"])
    pred_matrix = []
    model_columns = []
    manifest_rows = []
    oof_long_rows = []

    for repeat in range(1, n_repeats + 1):
        kfold = KFold(n_splits=n_splits, shuffle=True, random_state=int(config["base_seed"]) + repeat)
        print(f"Training repeated 10-fold CV seed ensemble repeat {repeat}/{n_repeats}")
        for fold, (train_idx, test_idx) in enumerate(kfold.split(x_exp), start=1):
            x_train = x_exp.iloc[train_idx].reset_index(drop=True)
            y_train = y_exp.iloc[train_idx].reset_index(drop=True)
            x_test = x_exp.iloc[test_idx].reset_index(drop=True)
            y_test = y_exp.iloc[test_idx].reset_index(drop=True)
            y_true = y_test.iloc[:, 0].to_numpy(dtype=float)
            print(f"  Fold {fold}/{n_splits}")
            for seed_index in range(1, seed_per_fold + 1):
                seed = int(config["base_seed"]) + repeat * 100000 + fold * 1000 + seed_index
                model_id = f"repeat_{repeat:02d}_fold_{fold:02d}_seed_{seed_index:02d}"
                model_path = cv_dir / f"repeat_{repeat:03d}" / f"fold_{fold:03d}" / f"seed_{seed_index:03d}"
                model_path.mkdir(parents=True, exist_ok=True)
                print(f"    Training {model_id} with seed={seed}")
                premodel, pre_data, pre_history, pre_metric = _fit_premodel(seed, dict(config["pretrain"]))
                fitted = _fit_descriptor_to_z(
                    premodel,
                    pre_data["norm_x"],
                    x_train,
                    y_train,
                    x_test,
                    y_test,
                    dict(config["train"]),
                )
                grid_pred = predict_member_grid(
                    premodel,
                    fitted["model"],
                    pre_data["norm_x"],
                    fitted["norm_y"],
                    grid,
                    config["prediction_chunk_size"],
                )
                pred_matrix.append(grid_pred.astype(np.float32))
                model_columns.append(model_id)
                test_pred = fitted["eval_pred"].reshape(-1)

                torch.save(premodel.state_dict(), model_path / "pretrain_model_weights.pth")
                torch.save(fitted["model"].state_dict(), model_path / "train_model_weights.pth")
                save_pickle(pre_data["norm_x"], model_path / "pretrain_norm_x.pkl")
                save_pickle(pre_data["norm_y"], model_path / "pretrain_norm_y.pkl")
                save_pickle(fitted["norm_y"], model_path / "train_norm_y.pkl")
                save_json(
                    {
                        "model_id": model_id,
                        "repeat": repeat,
                        "fold": fold,
                        "seed_index": seed_index,
                        "seed": seed,
                        "train_rows": int(len(train_idx)),
                        "test_rows": int(len(test_idx)),
                        "pretrain_metrics": pre_metric,
                        "train_metrics": fitted["metric"],
                        "pretrain_history": pre_history,
                        "train_history": fitted["history"],
                    },
                    model_path / "model_info.json",
                )
                manifest_rows.append(
                    {
                        "model_id": model_id,
                        "repeat": repeat,
                        "fold": fold,
                        "seed_index": seed_index,
                        "seed": seed,
                        "train_rows": int(len(train_idx)),
                        "test_rows": int(len(test_idx)),
                        "model_dir": model_path.relative_to(REPO_ROOT).as_posix(),
                    }
                )
                for local_i, original_idx in enumerate(test_idx):
                    oof_long_rows.append(
                        {
                            "row_index": int(original_idx),
                            "repeat": repeat,
                            "fold": fold,
                            "seed_index": seed_index,
                            "model_id": model_id,
                            "y_true": float(y_true[local_i]),
                            "pred_oof": float(test_pred[local_i]),
                        }
                    )

    pred_matrix = np.vstack(pred_matrix).T
    grid_summary = grid.copy()
    grid_summary["sigma_ML_epi"] = pred_matrix.std(axis=1, ddof=0)
    grid_pred_df = pd.concat([grid.copy(), pd.DataFrame(pred_matrix, columns=model_columns)], axis=1)

    grid_prediction_path = save_large_table(
        grid_pred_df,
        result_dir / "grid_cv_seed_predictions",
        preferred_format=config["grid_prediction_format"],
    )
    grid_summary.to_excel(result_dir / "grid_cv_seed_summary.xlsx", index=False)
    oof_long = pd.DataFrame(oof_long_rows).sort_values(["row_index", "repeat", "fold", "seed_index"])
    oof_long.to_excel(result_dir / "cv_seed_oof_predictions_long.xlsx", index=False)
    oof_rows = []
    for row_index, group in oof_long.groupby("row_index", sort=True):
        y_true = float(group["y_true"].iloc[0])
        oof_prediction_center = float(group["pred_oof"].mean())
        sigma_oof = float(group["pred_oof"].std(ddof=0))
        residual = float(y_true - oof_prediction_center)
        oof_rows.append(
            {
                "row_index": int(row_index),
                "y_true": y_true,
                "sigma_ML_epi_oof": sigma_oof,
                "oof_prediction_count": int(len(group)),
                "oof_repeat_count": int(group["repeat"].nunique()),
                "residual": residual,
                "abs_residual": float(abs(residual)),
                "residual_var_proxy_ML_only": float(max(residual**2 - sigma_oof**2, 0.0)),
            }
        )
    pd.DataFrame(oof_rows).to_excel(result_dir / "cv_seed_oof_predictions.xlsx", index=False)
    pd.DataFrame(manifest_rows).to_excel(result_dir / "cv_seed_model_manifest.xlsx", index=False)
    save_json(
        {
            "run_id": config["run_id"],
            "n_repeats": n_repeats,
            "n_splits": n_splits,
            "seed_per_fold": seed_per_fold,
            "total_models": int(n_repeats * n_splits * seed_per_fold),
            "oof_predictions_per_experiment": int(n_repeats * seed_per_fold),
            "grid_rows": int(len(grid)),
            "grid_prediction_path": grid_prediction_path.relative_to(REPO_ROOT).as_posix(),
            "method": "Repeated 10-fold CV seed ensemble; no full-data seed ensemble is trained.",
            "config": config,
        },
        result_dir / "cv_seed_ensemble_summary.json",
    )
    save_json(config, model_dir / "pipeline_config.json")
    return result_dir


def predict_member_grid(premodel, model, norm_x, norm_y, x_frame, chunk_size):
    pred_parts = []
    for start in range(0, len(x_frame), int(chunk_size)):
        chunk = x_frame.iloc[start : start + int(chunk_size)][RATIO_COLUMNS]
        x_scaled = norm_x.transform(chunk)
        with torch.no_grad():
            descriptor = premodel(torch.tensor(x_scaled).to(torch.float32))
            pred_scaled = model(descriptor).detach().numpy()
        pred_parts.append(norm_y.inverse_transform(pred_scaled).reshape(-1))
    return np.concatenate(pred_parts)


def save_large_table(frame, path_without_suffix, preferred_format="parquet"):
    path_without_suffix = Path(path_without_suffix)
    if preferred_format == "parquet":
        try:
            path = path_without_suffix.with_suffix(".parquet")
            frame.to_parquet(path, index=False)
            return path
        except Exception:
            pass
    path = path_without_suffix.with_suffix(".xlsx")
    frame.to_excel(path, index=False)
    return path


def read_large_table(path_without_suffix):
    path_without_suffix = Path(path_without_suffix)
    parquet_path = path_without_suffix.with_suffix(".parquet")
    excel_path = path_without_suffix.with_suffix(".xlsx")
    if parquet_path.exists():
        return pd.read_parquet(parquet_path)
    if excel_path.exists():
        return pd.read_excel(excel_path)
    raise FileNotFoundError(f"Cannot find {parquet_path} or {excel_path}")


def euclidean_distance_matrix(a, b):
    a = np.asarray(a, dtype=float)
    b = np.asarray(b, dtype=float)
    return np.sqrt(((a[:, None, :] - b[None, :, :]) ** 2).sum(axis=2))


def compositions_to_unit(values):
    values = np.asarray(values, dtype=float)
    row_sum = values.sum(axis=1, keepdims=True)
    if np.any(row_sum <= 0):
        raise ValueError("Composition rows must have positive sums before ILR transform.")
    return values / row_sum


def multiplicative_zero_replacement(values, delta):
    comp = compositions_to_unit(values)
    out = comp.copy()
    delta = float(delta)
    for row_index in range(out.shape[0]):
        row = out[row_index].copy()
        zero_mask = row <= 0
        zero_count = int(zero_mask.sum())
        if zero_count == 0:
            continue
        replacement_total = zero_count * delta
        if replacement_total >= 1.0:
            raise ValueError("ilr_zero_replacement_delta is too large.")
        nonzero_sum = float(row[~zero_mask].sum())
        row[zero_mask] = delta
        row[~zero_mask] = row[~zero_mask] * ((1.0 - replacement_total) / nonzero_sum)
        out[row_index] = row
    return out


def ilr_basis(dim):
    basis = np.zeros((dim - 1, dim), dtype=float)
    for idx in range(1, dim):
        basis[idx - 1, :idx] = 1.0 / np.sqrt(idx * (idx + 1))
        basis[idx - 1, idx] = -idx / np.sqrt(idx * (idx + 1))
    return basis


def ilr_transform(values, delta):
    comp = multiplicative_zero_replacement(values, delta)
    return np.log(comp) @ ilr_basis(comp.shape[1]).T


def knn_mean_distances(query_values, train_values, k_neighbors, chunk_size=50000, exclude_self=False):
    query_values = np.asarray(query_values, dtype=float)
    train_values = np.asarray(train_values, dtype=float)
    k_neighbors = max(1, min(int(k_neighbors), len(train_values) - int(exclude_self)))
    parts = []
    same_array = exclude_self and len(query_values) == len(train_values) and np.array_equal(query_values, train_values)
    for start in range(0, len(query_values), int(chunk_size)):
        query_chunk = query_values[start : start + int(chunk_size)]
        distances = euclidean_distance_matrix(query_chunk, train_values)
        if same_array:
            row_indices = np.arange(start, min(start + len(query_chunk), len(train_values)))
            distances[np.arange(len(row_indices)), row_indices] = np.inf
        nearest = np.partition(distances, k_neighbors - 1, axis=1)[:, :k_neighbors]
        parts.append(nearest.mean(axis=1))
    return np.concatenate(parts)


def add_ilr_ad95(table, config):
    experiment_x, _ = load_xy("experiment.xlsx", "metals", "potential")
    exp_ilr = ilr_transform(experiment_x[RATIO_COLUMNS].to_numpy(float), config["ilr_zero_replacement_delta"])
    grid_ilr = ilr_transform(table[RATIO_COLUMNS].to_numpy(float), config["ilr_zero_replacement_delta"])
    train_dist = knn_mean_distances(
        exp_ilr, exp_ilr, config["ad_k_neighbors"], config["prediction_chunk_size"], exclude_self=True
    )
    d_ad = float(np.quantile(train_dist, config["ad_quantile"]))
    grid_dist = knn_mean_distances(
        grid_ilr, exp_ilr, config["ad_k_neighbors"], config["prediction_chunk_size"], exclude_self=False
    )
    out = table.copy()
    out["dKNN_ilr"] = grid_dist
    out["dAD95_ilr"] = d_ad
    out["normalized_AD_ilr"] = grid_dist / d_ad
    out["AD95_ilr"] = np.where(grid_dist <= d_ad, AD95_LABEL_INSIDE, AD95_LABEL_OUTSIDE)
    return out, {
        "method": "ILR-5NN AD95",
        "ad_k_neighbors": int(config["ad_k_neighbors"]),
        "ad_quantile": float(config["ad_quantile"]),
        "dAD95_ilr": d_ad,
    }


def build_applicability_domain(config):
    _, result_dir = run_dirs(config)
    grid_summary = pd.read_excel(result_dir / "grid_cv_seed_summary.xlsx")
    grid_ad, ad_summary = add_ilr_ad95(grid_summary, config)
    grid_ad.to_excel(result_dir / "applicability_domain_grid.xlsx", index=False)

    experiment_x, experiment_y = load_xy("experiment.xlsx", "metals", "potential")
    exp_table = experiment_x[RATIO_COLUMNS].copy()
    exp_table["row_index"] = np.arange(len(exp_table))
    exp_table["Exp_ID"] = [f"Exp-{idx}" for idx in range(1, len(exp_table) + 1)]
    exp_table["E10"] = experiment_y.iloc[:, 0].to_numpy(dtype=float)
    exp_ilr = ilr_transform(exp_table[RATIO_COLUMNS].to_numpy(float), config["ilr_zero_replacement_delta"])
    exp_dist = knn_mean_distances(
        exp_ilr,
        exp_ilr,
        config["ad_k_neighbors"],
        config["prediction_chunk_size"],
        exclude_self=True,
    )
    exp_table["dKNN_ilr"] = exp_dist
    exp_table["dAD95_ilr"] = ad_summary["dAD95_ilr"]
    exp_table["normalized_AD_ilr"] = exp_dist / ad_summary["dAD95_ilr"]
    exp_table["AD95_ilr"] = np.where(exp_dist <= ad_summary["dAD95_ilr"], AD95_LABEL_INSIDE, AD95_LABEL_OUTSIDE)
    exp_table.to_excel(result_dir / "applicability_domain_experiment.xlsx", index=False)
    save_json(
        {
            "run_id": config["run_id"],
            "grid_rows": int(len(grid_ad)),
            "experiment_rows": int(len(exp_table)),
            "ad_summary": ad_summary,
            "grid_AD95_counts": {
                str(key): int(value) for key, value in grid_ad["AD95_ilr"].value_counts().to_dict().items()
            },
            "experiment_AD95_counts": {
                str(key): int(value) for key, value in exp_table["AD95_ilr"].value_counts().to_dict().items()
            },
            "config": config,
        },
        result_dir / "applicability_domain_summary.json",
    )
    return result_dir


def sigma_distribution_stats(values):
    series = pd.Series(values, dtype=float).dropna()
    if series.empty:
        return {
            "count": 0,
            "min": np.nan,
            "Q25": np.nan,
            "median": np.nan,
            "Q75": np.nan,
            "Q90": np.nan,
            "Q95": np.nan,
            "max": np.nan,
        }
    return {
        "count": int(len(series)),
        "min": float(series.min()),
        "Q25": float(series.quantile(0.25)),
        "median": float(series.quantile(0.50)),
        "Q75": float(series.quantile(0.75)),
        "Q90": float(series.quantile(0.90)),
        "Q95": float(series.quantile(0.95)),
        "max": float(series.max()),
    }


def build_sigma_threshold_transferability_diagnostics(result_dir, merged_oof, tau_epi):
    grid_sigma = pd.read_excel(result_dir / "grid_cv_seed_summary.xlsx")
    grid_ad = pd.read_excel(result_dir / "applicability_domain_grid.xlsx")
    grid = merge_on_ratio(grid_sigma, grid_ad, ["AD95_ilr"])

    oof_ad_in = merged_oof.loc[merged_oof["AD95_ilr"] == AD95_LABEL_INSIDE].copy()
    grid_ad_in = grid.loc[grid["AD95_ilr"] == AD95_LABEL_INSIDE].copy()
    grid_ad_out = grid.loc[grid["AD95_ilr"] == AD95_LABEL_OUTSIDE].copy()

    oof_stats = sigma_distribution_stats(oof_ad_in["sigma_ML_epi_oof"])
    grid_ad_in_stats = sigma_distribution_stats(grid_ad_in["sigma_ML_epi"])
    grid_ad_out_stats = sigma_distribution_stats(grid_ad_out["sigma_ML_epi"])

    median_oof = oof_stats["median"]
    median_grid_ad_in = grid_ad_in_stats["median"]
    median_ratio = (
        float(median_grid_ad_in / median_oof)
        if pd.notna(median_oof) and median_oof != 0 and pd.notna(median_grid_ad_in)
        else np.nan
    )
    diagnostics = {
        "tau_epi": float(tau_epi),
        "oof_AD_in_sigma_ML_epi_oof": oof_stats,
        "grid_AD_in_sigma_ML_epi": grid_ad_in_stats,
        "grid_AD_out_sigma_ML_epi": grid_ad_out_stats,
        "fraction_grid_sigma_le_tau_epi": float((grid["sigma_ML_epi"] <= tau_epi).mean()),
        "fraction_AD_in_grid_sigma_le_tau_epi": float((grid_ad_in["sigma_ML_epi"] <= tau_epi).mean())
        if len(grid_ad_in)
        else np.nan,
        "fraction_AD_out_grid_sigma_le_tau_epi": float((grid_ad_out["sigma_ML_epi"] <= tau_epi).mean())
        if len(grid_ad_out)
        else np.nan,
        "min_grid_sigma_AD_in_gt_tau_epi": bool(grid_ad_in_stats["min"] > tau_epi)
        if pd.notna(grid_ad_in_stats["min"])
        else None,
        "median_grid_AD_in_over_median_oof_AD_in": median_ratio,
    }
    rows = []
    for group_name, stats in [
        ("OOF_AD_in_sigma_ML_epi_oof", oof_stats),
        ("grid_AD_in_sigma_ML_epi", grid_ad_in_stats),
        ("grid_AD_out_sigma_ML_epi", grid_ad_out_stats),
    ]:
        row = {"group": group_name}
        row.update(stats)
        rows.append(row)
    fractions = pd.DataFrame(
        [
            {"item": key, "value": value}
            for key, value in diagnostics.items()
            if key
            not in {
                "oof_AD_in_sigma_ML_epi_oof",
                "grid_AD_in_sigma_ML_epi",
                "grid_AD_out_sigma_ML_epi",
            }
        ]
    )
    with pd.ExcelWriter(result_dir / "sigma_threshold_transferability_diagnostics.xlsx") as writer:
        pd.DataFrame(rows).to_excel(writer, sheet_name="sigma_distributions", index=False)
        fractions.to_excel(writer, sheet_name="threshold_position", index=False)
    save_json(diagnostics, result_dir / "sigma_threshold_transferability_diagnostics.json")
    return diagnostics


def build_oof_risk_coverage_calibration(config):
    _, result_dir = run_dirs(config)
    oof = pd.read_excel(result_dir / "cv_seed_oof_predictions.xlsx")
    exp_ad = pd.read_excel(result_dir / "applicability_domain_experiment.xlsx")
    merged = oof.merge(
        exp_ad[["row_index", "Exp_ID", "E10", "dKNN_ilr", "dAD95_ilr", "normalized_AD_ilr", "AD95_ilr"]],
        on="row_index",
        how="left",
    )
    if merged["AD95_ilr"].isna().any():
        raise ValueError("Some OOF rows could not be matched to applicability_domain_experiment.xlsx.")
    merged.to_excel(result_dir / "oof_with_ad_and_risk_inputs.xlsx", index=False)

    accepted_base = merged.loc[merged["AD95_ilr"] == AD95_LABEL_INSIDE].copy()
    if accepted_base.empty:
        raise ValueError("No AD-in OOF samples are available for risk-coverage calibration.")
    thresholds = np.sort(accepted_base["sigma_ML_epi_oof"].unique())
    rows = []
    n_ad_in = len(accepted_base)
    for threshold in thresholds:
        accepted = accepted_base.loc[accepted_base["sigma_ML_epi_oof"] <= threshold]
        if accepted.empty:
            continue
        risk = float(accepted["abs_residual"].mean())
        se = float(accepted["abs_residual"].std(ddof=1) / np.sqrt(len(accepted))) if len(accepted) > 1 else 0.0
        rows.append(
            {
                "threshold": float(threshold),
                "accepted_count": int(len(accepted)),
                "N_AD_in": int(n_ad_in),
                "coverage": float(len(accepted) / n_ad_in),
                "risk_mae": risk,
                "risk_se": se,
            }
        )
    curve = pd.DataFrame(rows)
    if curve.empty:
        raise ValueError("Risk-coverage curve is empty.")
    min_idx = curve["risk_mae"].idxmin()
    risk_min = float(curve.loc[min_idx, "risk_mae"])
    se_min = float(curve.loc[min_idx, "risk_se"])
    threshold_min = float(curve.loc[min_idx, "threshold"])
    accepted_count_at_threshold_min = int(curve.loc[min_idx, "accepted_count"])
    coverage_at_threshold_min = float(curve.loc[min_idx, "coverage"])
    eligible = curve.loc[curve["risk_mae"] <= risk_min + se_min]
    tau_epi = float(eligible["threshold"].max())
    curve["within_one_se_of_min"] = curve["risk_mae"] <= risk_min + se_min
    curve["selected_tau_epi"] = tau_epi
    curve.to_excel(result_dir / "oof_risk_coverage_curve.xlsx", index=False)
    tau_row = curve.loc[curve["threshold"] == tau_epi].iloc[-1]
    risk_at_tau = float(tau_row["risk_mae"])
    coverage_at_tau = float(tau_row["coverage"])
    risk_all_ad_in = float(accepted_base["abs_residual"].mean())
    sigma_transfer = build_sigma_threshold_transferability_diagnostics(result_dir, merged, tau_epi)
    summary = {
        "run_id": config["run_id"],
        "method": "OOF risk-coverage calibration within AD-in samples using repeated CV seed ensemble sigma.",
        "N_total_oof": int(len(merged)),
        "N_AD_in": int(n_ad_in),
        "oof_predictions_per_experiment": int(merged["oof_prediction_count"].iloc[0])
        if "oof_prediction_count" in merged.columns and not merged.empty
        else None,
        "threshold_min": threshold_min,
        "accepted_count_at_threshold_min": accepted_count_at_threshold_min,
        "coverage_at_threshold_min": coverage_at_threshold_min,
        "risk_at_threshold_min": risk_min,
        "risk_min": risk_min,
        "SE_min": se_min,
        "tau_epi": tau_epi,
        "coverage_at_tau_epi": coverage_at_tau,
        "risk_at_tau_epi": risk_at_tau,
        "risk_all_AD_in": risk_all_ad_in,
        "risk_at_tau_minus_risk_min": float(risk_at_tau - risk_min),
        "risk_at_tau_minus_risk_min_mV": float((risk_at_tau - risk_min) * 1000.0),
        "risk_all_AD_in_minus_risk_min": float(risk_all_ad_in - risk_min),
        "risk_all_AD_in_minus_risk_min_mV": float((risk_all_ad_in - risk_min) * 1000.0),
        "risk_all_AD_in_within_one_SE": bool(risk_all_ad_in <= risk_min + se_min),
        "SE_min_over_risk_min": float(se_min / risk_min) if risk_min != 0 else np.nan,
        "sigma_threshold_transferability": sigma_transfer,
        "config": config,
    }
    save_json(summary, result_dir / "epistemic_threshold_summary.json")
    return result_dir


def analysis_predictions(config):
    path = resolve_repo_path(config["analysis_path"])
    data = pd.read_excel(path)
    ratio = data.iloc[:, :5].copy()
    ratio.columns = RATIO_COLUMNS
    value_col = "potential" if "potential" in data.columns else data.columns[-1]
    out = ratio.copy()
    out["predicted_potential"] = data[value_col].to_numpy(dtype=float)
    return out


def merge_on_ratio(left, right, columns):
    out = left.copy()
    merge_frame = right[RATIO_COLUMNS + columns].copy()
    for col in RATIO_COLUMNS:
        out[col] = out[col].round(8)
        merge_frame[col] = merge_frame[col].round(8)
    return out.merge(merge_frame, on=RATIO_COLUMNS, how="left")


def merge_central_predictions(table, config):
    left = table.copy()
    right = analysis_predictions(config)
    for col in RATIO_COLUMNS:
        left[col] = left[col].round(8)
        right[col] = right[col].round(8)
    return left.merge(right, on=RATIO_COLUMNS, how="left")


def ad_flag_from_label(value):
    if value == AD95_LABEL_INSIDE:
        return "AD-in"
    if value == AD95_LABEL_OUTSIDE:
        return "AD-out"
    return np.nan


def zone_label(ad_flag, u_flag):
    if ad_flag == "AD-in" and u_flag == "U-low":
        return "Z1"
    if ad_flag == "AD-out" and u_flag == "U-low":
        return "Z2"
    if ad_flag == "AD-in" and u_flag == "U-high":
        return "Z3"
    if ad_flag == "AD-out" and u_flag == "U-high":
        return "Z4"
    return np.nan


def overview_threshold_rows(config, result_dir):
    rows = [
        {"item": "run_id", "value": config["run_id"], "description": "Manual run identifier."},
        {
            "item": "cv_n_repeats",
            "value": config.get("cv_n_repeats", 1),
            "description": "Number of repeated CV partition rounds.",
        },
        {"item": "cv_n_splits", "value": config["cv_n_splits"], "description": "Number of CV folds."},
        {"item": "seed_per_fold", "value": config["seed_per_fold"], "description": "Seed models trained per CV fold."},
        {
            "item": "total_cv_seed_models",
            "value": int(config.get("cv_n_repeats", 1)) * int(config["cv_n_splits"]) * int(config["seed_per_fold"]),
            "description": "Total models used for grid sigma_ML_epi.",
        },
        {
            "item": "oof_predictions_per_experiment",
            "value": int(config.get("cv_n_repeats", 1)) * int(config["seed_per_fold"]),
            "description": "OOF predictions used to compute each experimental sigma_ML_epi_oof.",
        },
        {"item": "ad_k_neighbors", "value": config["ad_k_neighbors"], "description": "k in ILR-kNN AD."},
        {"item": "ad_quantile", "value": config["ad_quantile"], "description": "Quantile used for AD threshold."},
        {
            "item": "Zone_Z1",
            "value": "AD-in + U-low",
            "description": "Inside ILR-5NN AD95 and sigma_ML_epi <= tau_epi.",
        },
        {
            "item": "Zone_Z2",
            "value": "AD-out + U-low",
            "description": "Outside ILR-5NN AD95 and sigma_ML_epi <= tau_epi.",
        },
        {
            "item": "Zone_Z3",
            "value": "AD-in + U-high",
            "description": "Inside ILR-5NN AD95 and sigma_ML_epi > tau_epi.",
        },
        {
            "item": "Zone_Z4",
            "value": "AD-out + U-high",
            "description": "Outside ILR-5NN AD95 and sigma_ML_epi > tau_epi.",
        },
    ]
    ad_summary_path = result_dir / "applicability_domain_summary.json"
    if ad_summary_path.exists():
        with open(ad_summary_path, encoding="utf-8") as f:
            summary = json.load(f)
        ad_summary = summary.get("ad_summary", {})
        for key in ["dAD95_ilr", "method", "ad_k_neighbors", "ad_quantile"]:
            if key in ad_summary:
                rows.append({"item": key, "value": ad_summary[key], "description": "ILR-kNN AD95 parameter/result."})

    epi_summary_path = result_dir / "epistemic_threshold_summary.json"
    if epi_summary_path.exists():
        with open(epi_summary_path, encoding="utf-8") as f:
            summary = json.load(f)
        for key in [
            "tau_epi",
            "threshold_min",
            "accepted_count_at_threshold_min",
            "coverage_at_threshold_min",
            "risk_at_threshold_min",
            "risk_min",
            "SE_min",
            "coverage_at_tau_epi",
            "risk_at_tau_epi",
            "risk_all_AD_in",
            "risk_at_tau_minus_risk_min",
            "risk_at_tau_minus_risk_min_mV",
            "risk_all_AD_in_minus_risk_min",
            "risk_all_AD_in_minus_risk_min_mV",
            "risk_all_AD_in_within_one_SE",
            "SE_min_over_risk_min",
            "N_AD_in",
        ]:
            if key in summary:
                rows.append({"item": key, "value": summary[key], "description": "OOF risk-coverage calibration result."})
        sigma_transfer = summary.get("sigma_threshold_transferability", {})
        for key in [
            "fraction_grid_sigma_le_tau_epi",
            "fraction_AD_in_grid_sigma_le_tau_epi",
            "fraction_AD_out_grid_sigma_le_tau_epi",
            "min_grid_sigma_AD_in_gt_tau_epi",
            "median_grid_AD_in_over_median_oof_AD_in",
        ]:
            if key in sigma_transfer:
                rows.append(
                    {
                        "item": key,
                        "value": sigma_transfer[key],
                        "description": "OOF/grid sigma threshold transferability diagnostic.",
                    }
                )
    return pd.DataFrame(rows)


def write_grid_uncertainty_overview(config):
    _, result_dir = run_dirs(config)
    grid = analysis_predictions(config)
    grid = grid[RATIO_COLUMNS + ["predicted_potential"]]

    grid_summary_path = result_dir / "grid_cv_seed_summary.xlsx"
    if grid_summary_path.exists():
        grid_summary = pd.read_excel(grid_summary_path)
        grid = merge_on_ratio(grid, grid_summary, ["sigma_ML_epi"])

    ad_path = result_dir / "applicability_domain_grid.xlsx"
    if ad_path.exists():
        ad = pd.read_excel(ad_path)
        grid = merge_on_ratio(grid, ad, ["dKNN_ilr", "dAD95_ilr", "normalized_AD_ilr", "AD95_ilr"])
        grid["AD_flag"] = grid["AD95_ilr"].map(ad_flag_from_label)

    epi_summary_path = result_dir / "epistemic_threshold_summary.json"
    if epi_summary_path.exists() and "sigma_ML_epi" in grid.columns:
        with open(epi_summary_path, encoding="utf-8") as f:
            epi_summary = json.load(f)
        tau_epi = float(epi_summary["tau_epi"])
        grid["tau_epi"] = tau_epi
        grid["U_flag"] = np.where(grid["sigma_ML_epi"] <= tau_epi, "U-low", "U-high")
        grid["ML_epi_class"] = np.where(
            grid["sigma_ML_epi"] <= tau_epi,
            "low ML epistemic",
            "high ML epistemic",
        )
        if "AD_flag" in grid.columns:
            grid["Zone"] = [zone_label(ad, unc) for ad, unc in zip(grid["AD_flag"], grid["U_flag"])]
            grid["accepted_by_oof_calibration"] = grid["Zone"] == "Z1"

    preferred = RATIO_COLUMNS + [
        "predicted_potential",
        "sigma_ML_epi",
        "dKNN_ilr",
        "dAD95_ilr",
        "normalized_AD_ilr",
        "AD95_ilr",
        "AD_flag",
        "tau_epi",
        "U_flag",
        "ML_epi_class",
        "Zone",
        "accepted_by_oof_calibration",
    ]
    columns = [col for col in preferred if col in grid.columns]
    columns += [col for col in grid.columns if col not in columns]
    thresholds = overview_threshold_rows(config, result_dir)
    with pd.ExcelWriter(result_dir / "grid_uncertainty_overview.xlsx") as writer:
        grid[columns].to_excel(writer, sheet_name="grid", index=False)
        thresholds.to_excel(writer, sheet_name="thresholds", index=False)
    return result_dir / "grid_uncertainty_overview.xlsx"


def require_input_file(path, description, columns=None):
    if not path.exists():
        column_hint = f"\nRequired columns: {', '.join(columns)}" if columns else ""
        raise FileNotFoundError(
            f"Missing {description}: {path}\n"
            "Please provide this read-only source file before running the dependent pipeline step."
            f"{column_hint}"
        )


def load_repeatability_data(config):
    path = resolve_repo_path(config["repeatability_path"])
    require_input_file(path, "repeatability workbook")
    sheet_name = str(config.get("repeatability_sheet", "1"))
    raw = pd.read_excel(path, sheet_name=sheet_name)
    source_ratio_columns = ["Fe", "Co", "Ni", "Mn", "Zn"]
    missing = [col for col in source_ratio_columns if col not in raw.columns]
    if missing:
        raise ValueError(f"Repeatability sheet {sheet_name!r} is missing composition columns: {missing}")
    if raw.shape[1] < 10:
        raise ValueError(
            f"Repeatability sheet {sheet_name!r} must contain at least 10 columns; "
            "the final three columns must be replicate E10 measurements."
        )

    source_rep_columns = list(raw.columns[-3:])
    data = raw[source_ratio_columns + source_rep_columns].copy()
    data = data.dropna(subset=source_ratio_columns, how="all")
    if len(data) != int(config.get("expected_descriptor_representative_compositions", 5)):
        raise ValueError(
            f"Repeatability sheet {sheet_name!r} must contain exactly "
            f"{config.get('expected_descriptor_representative_compositions', 5)} composition rows; "
            f"found {len(data)}."
        )
    for col in source_ratio_columns + source_rep_columns:
        data[col] = pd.to_numeric(data[col], errors="coerce")
    if data[source_ratio_columns + source_rep_columns].isna().any().any():
        raise ValueError(
            f"Repeatability sheet {sheet_name!r} contains missing or non-numeric composition/replicate values."
        )
    ratio_sums = data[source_ratio_columns].sum(axis=1)
    if not np.allclose(ratio_sums, 1.0, atol=1e-6):
        raise ValueError("Repeatability composition fractions must sum to 1 for every row.")

    data = data.rename(
        columns={
            source_rep_columns[0]: "E10_rep1",
            source_rep_columns[1]: "E10_rep2",
            source_rep_columns[2]: "E10_rep3",
        }
    )
    return data[REQUIRED_REPEATABILITY_COLUMNS].reset_index(drop=True)


def pooled_repeatability_std(config):
    data = load_repeatability_data(config)
    rep_cols = ["E10_rep1", "E10_rep2", "E10_rep3"]
    values = data[rep_cols].to_numpy(dtype=float)
    group_mean = values.mean(axis=1, keepdims=True)
    numerator = float(((values - group_mean) ** 2).sum())
    denominator = int(values.shape[0] * (values.shape[1] - 1))
    return float(np.sqrt(numerator / denominator)), data


def local_smooth_from_experiments(table, experiment_values, config, value_column):
    experiment_x, _ = load_xy("experiment.xlsx", "metals", "potential")
    exp_ilr = ilr_transform(experiment_x[RATIO_COLUMNS].to_numpy(float), config["ilr_zero_replacement_delta"])
    grid_ilr = ilr_transform(table[RATIO_COLUMNS].to_numpy(float), config["ilr_zero_replacement_delta"])
    values = experiment_values.sort_values("row_index")[value_column].to_numpy(dtype=float)
    k_neighbors = max(1, min(int(config["residual_smoothing_k"]), len(exp_ilr)))
    parts = []
    for start in range(0, len(grid_ilr), int(config["prediction_chunk_size"])):
        chunk = grid_ilr[start : start + int(config["prediction_chunk_size"])]
        distances = euclidean_distance_matrix(chunk, exp_ilr)
        nearest_idx = np.argpartition(distances, k_neighbors - 1, axis=1)[:, :k_neighbors]
        nearest_dist = np.take_along_axis(distances, nearest_idx, axis=1)
        nearest_values = values[nearest_idx]
        weights = 1.0 / (nearest_dist + 1e-12)
        parts.append(np.sqrt((weights * nearest_values**2).sum(axis=1) / weights.sum(axis=1)))
    return np.concatenate(parts)


def prepare_empirical_residual(oof, repeatability_sd):
    out = oof.copy()
    out["residual_var_proxy"] = (
        out["abs_residual"] ** 2 - out["sigma_ML_epi_oof"] ** 2 - out["sigma_descriptor_oof"] ** 2
    ).clip(lower=0.0)
    out["sigma_residual_proxy"] = np.sqrt(out["residual_var_proxy"])
    out["sigma_residual_with_repeatability_floor"] = np.maximum(
        out["sigma_residual_proxy"], float(repeatability_sd)
    )
    out["repeatability_sd_floor"] = float(repeatability_sd)
    out["residual_proxy_floor_source"] = np.where(
        out["sigma_residual_proxy"] < float(repeatability_sd), "repeatability_floor", "OOF_residual_proxy"
    )
    return out


def load_cv_seed_models(config):
    model_dir, _ = run_dirs(config)
    manifest = pd.read_excel(run_dirs(config)[1] / "cv_seed_model_manifest.xlsx")
    models = []
    for _, row in manifest.iterrows():
        model_path = Path(row["model_dir"])
        pre_cfg = config["pretrain"]
        train_cfg = config["train"]
        premodel = DescriptorPreModel(
            pre_cfg["input_dim"], pre_cfg["hidden_units"], pre_cfg["output_dim"]
        )
        model = DescriptorToExperimentModel(
            train_cfg["input_dim"],
            train_cfg["hidden_units"],
            train_cfg["output_dim"],
            train_cfg["dropout"],
        )
        premodel.load_state_dict(torch.load(model_path / "pretrain_model_weights.pth", map_location="cpu"))
        model.load_state_dict(torch.load(model_path / "train_model_weights.pth", map_location="cpu"))
        premodel.eval()
        model.eval()
        models.append(
            {
                "model": model,
                "pretrain_norm_y": load_pickle(model_path / "pretrain_norm_y.pkl"),
                "train_norm_y": load_pickle(model_path / "train_norm_y.pkl"),
            }
        )
    return models


def require_columns(data, required, source):
    missing = [col for col in required if col not in data.columns]
    if missing:
        raise ValueError(f"{source} is missing columns: {missing}")


def parse_md_ratio(value):
    parts = str(value).strip().split(":")
    if len(parts) != 5:
        raise ValueError(f"Invalid Fe:Co:Ni:Mn:Zn ratio: {value!r}")
    values = np.asarray([float(item) for item in parts], dtype=float)
    if (values < 0).any() or values.sum() <= 0:
        raise ValueError(f"Invalid Fe:Co:Ni:Mn:Zn ratio: {value!r}")
    values = values / values.sum()
    by_element = dict(zip(["Fe", "Co", "Ni", "Mn", "Zn"], values))
    return {col: float(by_element[col]) for col in RATIO_COLUMNS}


def composition_key(frame):
    return {
        tuple(row)
        for row in frame[RATIO_COLUMNS].round(8).to_numpy(dtype=float)
    }


def read_temperature_descriptor_data(config):
    path = resolve_repo_path(config["md_descriptor_uncertainty_path"])
    require_input_file(path, "MD descriptor uncertainty workbook")
    sheet = config.get("md_temperature_sheet", "temperature")
    raw = pd.read_excel(path, sheet_name=sheet)
    required = [MD_RATIO_COLUMN, "Property", *MD_TEMPERATURE_COLUMNS, "\u03c3", "Max. |\u0394|"]
    require_columns(raw, required, f"{path.name}:{sheet}")
    raw = raw.copy()
    raw[MD_RATIO_COLUMN] = raw[MD_RATIO_COLUMN].ffill()
    raw = raw.dropna(subset=["Property"]).reset_index(drop=True)

    records = []
    for _, row in raw.iterrows():
        property_name = str(row["Property"]).strip()
        if property_name not in DFT_DESCRIPTOR_PROPERTY_MAP:
            raise ValueError(f"Unknown descriptor property {property_name!r} in {path.name}:{sheet}")
        condition_values = pd.to_numeric(row[MD_TEMPERATURE_COLUMNS], errors="coerce")
        sigma_value = pd.to_numeric(pd.Series([row["\u03c3"]]), errors="coerce").iloc[0]
        max_delta = pd.to_numeric(pd.Series([row["Max. |\u0394|"]]), errors="coerce").iloc[0]
        if condition_values.isna().any() or pd.isna(sigma_value) or pd.isna(max_delta):
            raise ValueError(f"Missing/non-numeric temperature descriptor data for {property_name!r}.")
        if sigma_value < 0 or max_delta < 0:
            raise ValueError(f"Temperature sigma and Max. |Delta| must be non-negative for {property_name!r}.")
        ratio_text = str(row[MD_RATIO_COLUMN]).strip()
        record = {
            **parse_md_ratio(ratio_text),
            "metal_ratio_Fe_Co_Ni_Mn_Zn": ratio_text,
            "Property": property_name,
            "descriptor": DFT_DESCRIPTOR_PROPERTY_MAP[property_name],
            **{col: float(condition_values[col]) for col in MD_TEMPERATURE_COLUMNS},
            "descriptor_center_physical": float(condition_values.mean()),
            "sigma_component": float(sigma_value),
            "max_abs_delta": float(max_delta),
            "source": "temperature",
            "representative_id": ratio_text,
            "source_workbook": path.name,
            "source_sheet": str(sheet),
            "sigma_definition": "reported sample sigma across temperature/time settings",
        }
        records.append(record)

    data = pd.DataFrame(records)
    expected = int(config.get("expected_descriptor_representative_compositions", 5))
    counts = data.groupby("metal_ratio_Fe_Co_Ni_Mn_Zn")["descriptor"].nunique()
    if len(counts) != expected or not (counts == len(DESCRIPTOR_COLUMNS)).all():
        raise ValueError(
            f"{path.name}:{sheet} must contain {expected} compositions x "
            f"{len(DESCRIPTOR_COLUMNS)} descriptors."
        )
    if data.duplicated(["metal_ratio_Fe_Co_Ni_Mn_Zn", "descriptor"]).any():
        raise ValueError(f"Duplicate composition/descriptor rows found in {path.name}:{sheet}.")
    return data


def read_force_field_descriptor_data(config):
    path = resolve_repo_path(config["md_descriptor_uncertainty_path"])
    sheet = config.get("md_force_field_sheet", "force_field")
    raw = pd.read_excel(path, sheet_name=sheet)
    required = [MD_RATIO_COLUMN, "Property", "UFF", "MACE", "|\u0394|"]
    require_columns(raw, required, f"{path.name}:{sheet}")
    raw = raw.copy()
    raw[MD_RATIO_COLUMN] = raw[MD_RATIO_COLUMN].ffill()
    raw = raw.dropna(subset=["Property"]).reset_index(drop=True)

    records = []
    for _, row in raw.iterrows():
        property_name = str(row["Property"]).strip()
        if property_name not in DFT_DESCRIPTOR_PROPERTY_MAP:
            raise ValueError(f"Unknown descriptor property {property_name!r} in {path.name}:{sheet}")
        values = pd.to_numeric(row[["UFF", "MACE", "|\u0394|"]], errors="coerce")
        if values.isna().any() or values["|\u0394|"] < 0:
            raise ValueError(f"Invalid force-field values for {property_name!r} in {path.name}:{sheet}.")
        ratio_text = str(row[MD_RATIO_COLUMN]).strip()
        delta = float(values["|\u0394|"])
        records.append(
            {
                **parse_md_ratio(ratio_text),
                "metal_ratio_Fe_Co_Ni_Mn_Zn": ratio_text,
                "Property": property_name,
                "descriptor": DFT_DESCRIPTOR_PROPERTY_MAP[property_name],
                "UFF": float(values["UFF"]),
                "MACE": float(values["MACE"]),
                "sigma_component": delta / np.sqrt(2.0),
                "max_abs_delta": delta,
                "source": "force_field",
                "representative_id": ratio_text,
                "source_workbook": path.name,
                "source_sheet": str(sheet),
                "sigma_definition": "two-setting sample SD = |Delta| / sqrt(2)",
            }
        )
    data = pd.DataFrame(records)
    if len(data) != len(DESCRIPTOR_COLUMNS) or data["descriptor"].nunique() != len(DESCRIPTOR_COLUMNS):
        raise ValueError(f"{path.name}:{sheet} must contain exactly one row for each descriptor.")
    return data


def read_single_site_descriptor_data(config):
    specs = [
        (
            "functional",
            config.get("single_site_functional_sheet", "sheet2_functional"),
            "\u03c3func",
            "\u0394PBEmax",
            False,
            "reported sample sigma across functionals",
        ),
        (
            "slab_size",
            config.get("single_site_slab_sheet", "sheet3_big_slab"),
            None,
            "|\u0394big-normal|",
            True,
            "two-setting sample SD = |Delta| / sqrt(2)",
        ),
        (
            "second_shell",
            config.get("single_site_second_shell_sheet", "sheet4_second_shell"),
            "\u03c3second_shell",
            "\u0394max",
            False,
            "reported sample sigma across second-shell elements",
        ),
    ]
    records = []
    for configured_path in config["single_site_descriptor_uncertainty_paths"]:
        path = resolve_repo_path(configured_path)
        require_input_file(path, "single-site descriptor uncertainty workbook")
        site_label = path.stem.removeprefix("single_site_results_")
        timestamp_parts = site_label.rsplit("_", 2)
        if len(timestamp_parts) == 3 and all(part.isdigit() for part in timestamp_parts[-2:]):
            site_label = timestamp_parts[0]
        available_sheets = set(pd.ExcelFile(path).sheet_names)
        required_sheets = {spec[1] for spec in specs}
        missing_sheets = sorted(required_sheets - available_sheets)
        if missing_sheets:
            raise ValueError(f"{path.name} is missing sheets: {missing_sheets}")
        for source, sheet, sigma_column, delta_column, derive_sigma, definition in specs:
            raw = pd.read_excel(path, sheet_name=sheet)
            required = ["Property", delta_column]
            if sigma_column:
                required.append(sigma_column)
            require_columns(raw, required, f"{path.name}:{sheet}")
            raw = raw.dropna(subset=["Property"]).reset_index(drop=True)
            for _, row in raw.iterrows():
                property_name = str(row["Property"]).strip()
                if property_name not in DFT_DESCRIPTOR_PROPERTY_MAP:
                    raise ValueError(f"Unknown descriptor property {property_name!r} in {path.name}:{sheet}")
                max_delta = pd.to_numeric(pd.Series([row[delta_column]]), errors="coerce").iloc[0]
                sigma_value = max_delta / np.sqrt(2.0) if derive_sigma else pd.to_numeric(
                    pd.Series([row[sigma_column]]), errors="coerce"
                ).iloc[0]
                if pd.isna(max_delta) or pd.isna(sigma_value) or max_delta < 0 or sigma_value < 0:
                    raise ValueError(f"Invalid uncertainty values for {property_name!r} in {path.name}:{sheet}.")
                records.append(
                    {
                        "site": site_label,
                        "source": source,
                        "Property": property_name,
                        "descriptor": DFT_DESCRIPTOR_PROPERTY_MAP[property_name],
                        "sigma_component": float(sigma_value),
                        "max_abs_delta": float(max_delta),
                        "representative_id": site_label,
                        "source_workbook": path.name,
                        "source_sheet": str(sheet),
                        "sigma_definition": definition,
                    }
                )
    data = pd.DataFrame(records)
    expected_per_source = len(config["single_site_descriptor_uncertainty_paths"]) * len(DESCRIPTOR_COLUMNS)
    counts = data.groupby("source").size()
    if set(counts.index) != {"functional", "slab_size", "second_shell"} or not (
        counts == expected_per_source
    ).all():
        raise ValueError("Single-site workbooks must contain all four descriptors for all three uncertainty sources.")
    return data


def build_descriptor_source_summary(temperature, force_field, single_site, config):
    if config.get("descriptor_representative_aggregation", "rms") != "rms":
        raise ValueError("descriptor_representative_aggregation currently supports only 'rms'.")
    if config.get("two_point_delta_to_sigma", "sample_sd") != "sample_sd":
        raise ValueError("two_point_delta_to_sigma currently supports only 'sample_sd'.")

    common = [
        "source",
        "representative_id",
        "Property",
        "descriptor",
        "sigma_component",
        "max_abs_delta",
        "source_workbook",
        "source_sheet",
        "sigma_definition",
    ]
    source_detail = pd.concat(
        [temperature[common], force_field[common], single_site[common]], ignore_index=True
    )
    if source_detail[["sigma_component", "max_abs_delta"]].isna().any().any():
        raise ValueError("DFT descriptor uncertainty inputs contain missing sigma or maximum-delta values.")
    if (source_detail[["sigma_component", "max_abs_delta"]] < 0).any().any():
        raise ValueError("DFT descriptor uncertainty inputs must be non-negative.")

    records = []
    for (source, descriptor), group in source_detail.groupby(["source", "descriptor"], sort=False):
        sigma_values = group["sigma_component"].to_numpy(dtype=float)
        delta_values = group["max_abs_delta"].to_numpy(dtype=float)
        records.append(
            {
                "source": source,
                "descriptor": descriptor,
                "representative_count": int(len(group)),
                "sigma_rms": float(np.sqrt(np.mean(sigma_values**2))),
                "sigma_max": float(np.max(sigma_values)),
                "max_abs_delta_rms": float(np.sqrt(np.mean(delta_values**2))),
                "max_abs_delta_max": float(np.max(delta_values)),
            }
        )
    source_summary = pd.DataFrame(records)
    expected_sources = {"temperature", "force_field", "functional", "slab_size", "second_shell"}
    for descriptor in DESCRIPTOR_COLUMNS:
        found = set(source_summary.loc[source_summary["descriptor"] == descriptor, "source"])
        if found != expected_sources:
            raise ValueError(f"Descriptor {descriptor!r} is missing uncertainty sources: {sorted(expected_sources - found)}")

    combined_records = []
    for descriptor in DESCRIPTOR_COLUMNS:
        rows = source_summary[source_summary["descriptor"] == descriptor]
        source_sigmas = dict(zip(rows["source"], rows["sigma_rms"]))
        combined_records.append(
            {
                "descriptor": descriptor,
                **{f"sigma_{source}": float(source_sigmas[source]) for source in sorted(expected_sources)},
                "combined_descriptor_sigma": float(np.sqrt(np.sum(np.square(list(source_sigmas.values()))))),
            }
        )
    return source_detail, source_summary, pd.DataFrame(combined_records)


def native_setting_records(data, source, group_id, setting_columns, source_workbook, source_sheet):
    frame = data.copy()
    if "descriptor" not in frame.columns:
        frame["descriptor"] = frame["Property"].astype(str).str.strip().map(DFT_DESCRIPTOR_PROPERTY_MAP)
    if frame["descriptor"].isna().any():
        unknown = frame.loc[frame["descriptor"].isna(), "Property"].astype(str).tolist()
        raise ValueError(f"Unknown descriptor properties in {source_workbook}:{source_sheet}: {unknown}")
    if frame.duplicated("descriptor").any() or set(frame["descriptor"]) != set(DESCRIPTOR_COLUMNS):
        raise ValueError(
            f"{source_workbook}:{source_sheet}, group {group_id!r}, must contain each descriptor exactly once."
        )
    by_descriptor = frame.set_index("descriptor")
    records = []
    for setting in setting_columns:
        values = pd.to_numeric(by_descriptor[setting], errors="coerce")
        if values.isna().any():
            raise ValueError(
                f"Missing/non-numeric setting values in {source_workbook}:{source_sheet}, "
                f"group {group_id!r}, setting {setting!r}."
            )
        records.append(
            {
                "source": source,
                "group_id": str(group_id),
                "setting": str(setting),
                **{descriptor: float(values.loc[descriptor]) for descriptor in DESCRIPTOR_COLUMNS},
                "source_workbook": source_workbook,
                "source_sheet": str(source_sheet),
            }
        )
    return records


def single_site_label(path):
    label = path.stem.removeprefix("single_site_results_")
    timestamp_parts = label.rsplit("_", 2)
    if len(timestamp_parts) == 3 and all(part.isdigit() for part in timestamp_parts[-2:]):
        return timestamp_parts[0]
    return label


def build_native_descriptor_settings(config, temperature, force_field):
    records = []
    md_path = resolve_repo_path(config["md_descriptor_uncertainty_path"])
    for ratio_text, group in temperature.groupby("metal_ratio_Fe_Co_Ni_Mn_Zn", sort=False):
        records.extend(
            native_setting_records(
                group,
                "temperature",
                ratio_text,
                MD_TEMPERATURE_COLUMNS,
                md_path.name,
                config.get("md_temperature_sheet", "temperature"),
            )
        )
    records.extend(
        native_setting_records(
            force_field,
            "force_field",
            force_field.iloc[0]["metal_ratio_Fe_Co_Ni_Mn_Zn"],
            ["UFF", "MACE"],
            md_path.name,
            config.get("md_force_field_sheet", "force_field"),
        )
    )

    single_site_specs = [
        ("functional", config.get("single_site_functional_sheet", "sheet2_functional"), ["PBE", "PBEsol", "RPBE", "r2SCAN"]),
        ("slab_size", config.get("single_site_slab_sheet", "sheet3_big_slab"), ["normal", "big_slab"]),
        ("second_shell", config.get("single_site_second_shell_sheet", "sheet4_second_shell"), ["Fe", "Co", "Ni", "Mn", "Zn"]),
    ]
    for configured_path in config["single_site_descriptor_uncertainty_paths"]:
        path = resolve_repo_path(configured_path)
        group_id = single_site_label(path)
        for source, sheet, settings in single_site_specs:
            raw = pd.read_excel(path, sheet_name=sheet).dropna(subset=["Property"]).reset_index(drop=True)
            require_columns(raw, ["Property", *settings], f"{path.name}:{sheet}")
            records.extend(
                native_setting_records(raw, source, group_id, settings, path.name, sheet)
            )
    settings = pd.DataFrame(records)
    expected_group_counts = {
        "functional": 2,
        "second_shell": 2,
        "slab_size": 2,
        "temperature": int(config.get("expected_descriptor_representative_compositions", 5)),
        "force_field": 1,
    }
    expected_setting_counts = {
        "functional": 4,
        "second_shell": 5,
        "slab_size": 2,
        "temperature": len(MD_TEMPERATURE_COLUMNS),
        "force_field": 2,
    }
    for source in DESCRIPTOR_UNCERTAINTY_SOURCES:
        source_rows = settings[settings["source"] == source]
        if source_rows["group_id"].nunique() != expected_group_counts[source]:
            raise ValueError(f"Unexpected group count for {source}: {source_rows['group_id'].nunique()}")
        counts = source_rows.groupby("group_id")["setting"].nunique()
        if not (counts == expected_setting_counts[source]).all():
            raise ValueError(f"Unexpected setting count for {source}: {counts.to_dict()}")
    return settings


def save_formatted_workbook(output_path, sheets):
    output_path = Path(output_path)
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, frame in sheets.items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
            worksheet = writer.sheets[sheet_name]
            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = worksheet.dimensions
            worksheet.row_dimensions[1].height = 24
            for cell in worksheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.alignment = Alignment(horizontal="center", vertical="center")
            wrapped_columns = []
            for column_index, column_name in enumerate(frame.columns):
                column = frame.iloc[:, column_index]
                cell_lengths = column.map(lambda value: len("" if pd.isna(value) else str(value)))
                maximum_length = max(len(str(column_name)), int(cell_lengths.max()) if len(frame) else 0)
                width = min(maximum_length + 2, 48)
                excel_column = get_column_letter(column_index + 1)
                worksheet.column_dimensions[excel_column].width = max(width, 12 if pd.api.types.is_numeric_dtype(column) else 10)
                if pd.api.types.is_numeric_dtype(column) and not pd.api.types.is_bool_dtype(column):
                    for cell in worksheet[excel_column][1:]:
                        cell.number_format = "0.000000"
                elif maximum_length > 46:
                    wrapped_columns.append((column_index, 46))
                    for cell in worksheet[excel_column][1:]:
                        cell.alignment = Alignment(vertical="top", wrap_text=True)
            for row_offset, row_values in enumerate(frame.itertuples(index=False, name=None), start=2):
                line_count = 1
                for column_index, characters_per_line in wrapped_columns:
                    value = "" if pd.isna(row_values[column_index]) else str(row_values[column_index])
                    line_count = max(line_count, int(np.ceil(len(value) / characters_per_line)))
                if line_count > 1:
                    worksheet.row_dimensions[row_offset].height = min(15 * line_count, 90)
    return output_path


def root_mean_square(values):
    array = np.asarray(values, dtype=float)
    if len(array) == 0:
        raise ValueError("Cannot calculate RMS from an empty array.")
    return float(np.sqrt(np.mean(array**2)))


def predict_native_descriptor_settings(settings, models, config):
    out = settings.copy()
    out["E10_prediction_V"] = predict_descriptor_samples(
        models, out[DESCRIPTOR_COLUMNS].to_numpy(dtype=float), config
    )
    return out


def native_input_scale_summary(settings):
    group_records = []
    for (source, group_id), group in settings.groupby(["source", "group_id"], sort=False):
        if len(group) < 2:
            raise ValueError(f"Source {source}, group {group_id!r}, needs at least two settings.")
        for descriptor in DESCRIPTOR_COLUMNS:
            group_records.append(
                {
                    "source": source,
                    "group_id": group_id,
                    "descriptor": descriptor,
                    "setting_count": int(len(group)),
                    "group_descriptor_sample_sd": float(group[descriptor].std(ddof=1)),
                }
            )
    group_detail = pd.DataFrame(group_records)
    records = []
    for source in DESCRIPTOR_SOURCE_REPORT_ORDER:
        source_rows = group_detail[group_detail["source"] == source]
        for descriptor in DESCRIPTOR_COLUMNS:
            rows = source_rows[source_rows["descriptor"] == descriptor]
            records.append(
                {
                    "source": source,
                    "descriptor": descriptor,
                    "perturbed_quantity": DESCRIPTOR_DISPLAY_NAMES[descriptor],
                    "scale_unit": DESCRIPTOR_UNITS[descriptor],
                    "group_count": int(len(rows)),
                    "setting_count_total": int(rows["setting_count"].sum()),
                    "perturbation_scale_RMS": root_mean_square(
                        rows["group_descriptor_sample_sd"]
                    ),
                }
            )
    return group_detail, pd.DataFrame(records)


def native_source_contributions(setting_predictions):
    group_records = []
    for (source, group_id), group in setting_predictions.groupby(["source", "group_id"], sort=False):
        predictions = group["E10_prediction_V"].to_numpy(dtype=float)
        group_records.append(
            {
                "source": source,
                "group_id": group_id,
                "setting_count": int(len(group)),
                "E10_mean_V": float(np.mean(predictions)),
                "E10_sample_sd_V": float(np.std(predictions, ddof=1)),
                "E10_min_V": float(np.min(predictions)),
                "E10_max_V": float(np.max(predictions)),
                "max_abs_E10_change_V": float(np.max(predictions) - np.min(predictions)),
            }
        )
    group_detail = pd.DataFrame(group_records)
    source_records = []
    for source in DESCRIPTOR_SOURCE_REPORT_ORDER:
        rows = group_detail[group_detail["source"] == source]
        uncertainty = root_mean_square(rows["E10_sample_sd_V"])
        source_records.append(
            {
                "source": source,
                "source_display": DESCRIPTOR_SOURCE_DISPLAY_NAMES[source],
                "perturbed_quantity": "All four descriptors jointly",
                "group_count": int(len(rows)),
                "setting_count_total": int(
                    setting_predictions.loc[
                        setting_predictions["source"] == source, ["group_id", "setting"]
                    ].drop_duplicates().shape[0]
                ),
                "propagated_E10_sigma_V": uncertainty,
                "propagated_E10_variance_V2": uncertainty**2,
            }
        )
    source_contribution = pd.DataFrame(source_records)
    variance_sum = float(source_contribution["propagated_E10_variance_V2"].sum())
    source_contribution["relative_contribution_to_sigma_desc_variance_pct"] = (
        100.0 * source_contribution["propagated_E10_variance_V2"] / variance_sum
    )
    return group_detail, source_contribution, float(np.sqrt(variance_sum))


def native_component_contributions(settings, models, config):
    group_records = []
    for (source, group_id), group in settings.groupby(["source", "group_id"], sort=False):
        center = group[DESCRIPTOR_COLUMNS].mean(axis=0).to_numpy(dtype=float)
        for descriptor_index, descriptor in enumerate(DESCRIPTOR_COLUMNS):
            descriptor_values = group[descriptor].to_numpy(dtype=float)
            scenarios = np.repeat(center.reshape(1, -1), len(group), axis=0)
            scenarios[:, descriptor_index] = descriptor_values
            predictions = predict_descriptor_samples(models, scenarios, config)
            group_records.append(
                {
                    "source": source,
                    "group_id": group_id,
                    "descriptor": descriptor,
                    "setting_count": int(len(group)),
                    "input_sample_sd": float(np.std(descriptor_values, ddof=1)),
                    "propagated_E10_sample_sd_V": float(np.std(predictions, ddof=1)),
                }
            )
    group_detail = pd.DataFrame(group_records)
    records = []
    for source in DESCRIPTOR_SOURCE_REPORT_ORDER:
        for descriptor in DESCRIPTOR_COLUMNS:
            rows = group_detail[
                (group_detail["source"] == source) & (group_detail["descriptor"] == descriptor)
            ]
            propagated = root_mean_square(rows["propagated_E10_sample_sd_V"])
            records.append(
                {
                    "source": source,
                    "descriptor": descriptor,
                    "perturbed_quantity": DESCRIPTOR_DISPLAY_NAMES[descriptor],
                    "perturbation_scale": root_mean_square(rows["input_sample_sd"]),
                    "scale_unit": DESCRIPTOR_UNITS[descriptor],
                    "propagated_E10_sigma_V": propagated,
                    "propagated_E10_variance_V2": propagated**2,
                    "usage": "diagnostic only; not included in sigma_descriptor",
                }
            )
    contribution = pd.DataFrame(records)
    diagnostic_variance_sum = float(contribution["propagated_E10_variance_V2"].sum())
    contribution["diagnostic_variance_share_pct"] = (
        100.0 * contribution["propagated_E10_variance_V2"] / diagnostic_variance_sum
    )
    return group_detail, contribution


def native_perturbation_scale_text(input_scale_summary, source):
    rows = input_scale_summary[input_scale_summary["source"] == source].set_index("descriptor")
    return "; ".join(
        f"{DESCRIPTOR_DISPLAY_NAMES[descriptor]}={rows.loc[descriptor, 'perturbation_scale_RMS']:.6f} "
        f"{DESCRIPTOR_UNITS[descriptor]}"
        for descriptor in DESCRIPTOR_COLUMNS
    )


def build_native_table_r3_6(source_contribution, input_scale_summary, sigma_descriptor):
    records = []
    for _, row in source_contribution.iterrows():
        source = row["source"]
        records.append(
            {
                "Uncertainty source": row["source_display"],
                "Perturbed quantity": "Four descriptors jointly",
                "Perturbation scale": native_perturbation_scale_text(input_scale_summary, source),
                "Propagated E10 uncertainty (V)": row["propagated_E10_sigma_V"],
                "Relative contribution to sigma_desc variance (%)": row[
                    "relative_contribution_to_sigma_desc_variance_pct"
                ],
            }
        )
    records.append(
        {
            "Uncertainty source": "Combined descriptor uncertainty",
            "Perturbed quantity": "Five uncertainty sources",
            "Perturbation scale": "Source-level E10 variances combined in quadrature",
            "Propagated E10 uncertainty (V)": sigma_descriptor,
            "Relative contribution to sigma_desc variance (%)": 100.0,
        }
    )
    return pd.DataFrame(records)


def build_descriptor_error_propagation(config):
    mode = config.get("descriptor_uncertainty_mode", "global")
    if mode != "global":
        raise ValueError(
            "The descriptor inputs currently support descriptor_uncertainty_mode='global' only."
        )
    temperature = read_temperature_descriptor_data(config)
    force_field = read_force_field_descriptor_data(config)
    single_site = read_single_site_descriptor_data(config)
    source_detail, _source_summary, _ = build_descriptor_source_summary(
        temperature, force_field, single_site, config
    )
    settings = build_native_descriptor_settings(config, temperature, force_field)
    models = load_cv_seed_models(config)
    setting_predictions = predict_native_descriptor_settings(settings, models, config)
    _input_scale_group_detail, input_scale_summary = native_input_scale_summary(settings)
    source_group_detail, source_contribution, sigma_descriptor = native_source_contributions(
        setting_predictions
    )
    _component_group_detail, component_contribution = native_component_contributions(
        settings, models, config
    )
    sensitivity = source_group_detail[
        ["source", "group_id", "setting_count", "E10_min_V", "E10_max_V", "max_abs_E10_change_V"]
    ].copy()
    sensitivity["included_in_sigma_descriptor"] = "No; absolute sensitivity audit only"
    table_r3_6 = build_native_table_r3_6(
        source_contribution, input_scale_summary, sigma_descriptor
    )
    _, result_dir = run_dirs(config)
    output_path = save_formatted_workbook(
        result_dir / "descriptor_error_propagation.xlsx",
        {
            "source_inputs": source_detail,
            "source_setting_predictions": setting_predictions,
            "source_group_detail": source_group_detail,
            "input_scale_summary": input_scale_summary,
            "component_contribution": component_contribution,
            "source_contribution": source_contribution,
            "max_sensitivity": sensitivity,
            "table_R3_6": table_r3_6,
        },
    )
    summary = {
        "mode": "global",
        "descriptor_input_space": config.get("descriptor_input_space", "physical"),
        "fixed_second_stage_model_count": int(len(models)),
        "propagation_method": "native complete descriptor setting vectors; no Monte Carlo perturbation",
        "setting_treatment": "all evaluated settings receive equal weight within each source group",
        "group_aggregation": "RMS of propagated E10 sample standard deviations",
        "two_setting_treatment": "two-setting E10 sample SD = absolute E10 difference / sqrt(2)",
        "source_sigma_combination": "quadrature of five source-level propagated E10 variances",
        "sigma_descriptor_global_V": sigma_descriptor,
        "maximum_delta_usage": "absolute sensitivity audit only; not added separately to sigma_total",
        "propagation": "global constant assigned to all grid and OOF points",
        "contribution_definition": "source propagated E10 variance divided by sigma_descriptor squared",
        "component_contribution_usage": "diagnostic only; component results are not additive and do not define sigma_descriptor",
        "direction_handling": (
            "signed slab-size and force-field E10 differences are retained for sensitivity reporting; "
            "only sample-SD magnitudes enter sigma_descriptor and no global mean correction is applied"
        ),
        "pretrain_RMSE_usage": "not added to descriptor uncertainty; first-stage training variability is represented within sigma_ML_epi",
        "source_paths": {
            "md": str(config["md_descriptor_uncertainty_path"]),
            "single_site": [str(path) for path in config["single_site_descriptor_uncertainty_paths"]],
        },
        "source_sheets": {
            "temperature": config.get("md_temperature_sheet", "temperature"),
            "force_field": config.get("md_force_field_sheet", "force_field"),
            "functional": config.get("single_site_functional_sheet", "sheet2_functional"),
            "slab_size": config.get("single_site_slab_sheet", "sheet3_big_slab"),
            "second_shell": config.get("single_site_second_shell_sheet", "sheet4_second_shell"),
        },
        "source_setting_counts": setting_predictions.groupby(["source", "group_id"])["setting"]
        .nunique()
        .reset_index(name="setting_count")
        .to_dict(orient="records"),
        "input_scale_summary": input_scale_summary.to_dict(orient="records"),
        "source_group_detail": source_group_detail.to_dict(orient="records"),
        "component_contribution": component_contribution.to_dict(orient="records"),
        "source_contribution": source_contribution.to_dict(orient="records"),
        "max_delta_sensitivity_by_source": sensitivity.groupby("source")["max_abs_E10_change_V"]
        .agg(["min", "median", "max"])
        .reset_index()
        .to_dict(orient="records"),
        "output_workbook": output_path.relative_to(REPO_ROOT).as_posix(),
    }
    save_json(summary, result_dir / "descriptor_uncertainty_summary.json")
    return result_dir


def predict_descriptor_samples(models, descriptors, config):
    descriptors = np.asarray(descriptors, dtype=float)
    chunk_size = max(1, int(config.get("descriptor_prediction_batch_size", len(descriptors))))
    prediction_parts = []
    for start in range(0, len(descriptors), chunk_size):
        descriptor_chunk = descriptors[start : start + chunk_size]
        model_preds = []
        with torch.no_grad():
            for item in models:
                if config.get("descriptor_input_space", "physical") == "physical":
                    descriptor_frame = pd.DataFrame(
                        descriptor_chunk, columns=SCALER_DESCRIPTOR_COLUMNS
                    )
                    expected_columns = list(
                        getattr(item["pretrain_norm_y"], "feature_names_in_", SCALER_DESCRIPTOR_COLUMNS)
                    )
                    if set(expected_columns) != set(SCALER_DESCRIPTOR_COLUMNS):
                        raise ValueError(
                            "pretrain_norm_y descriptor columns do not match the expected physical descriptor mapping: "
                            f"{expected_columns}"
                        )
                    descriptor_frame = descriptor_frame[expected_columns]
                    descriptor_input = item["pretrain_norm_y"].transform(descriptor_frame)
                elif config.get("descriptor_input_space") == "model_scaled":
                    descriptor_input = descriptor_chunk
                else:
                    raise ValueError("descriptor_input_space must be 'physical' or 'model_scaled'")
                descriptor_tensor = torch.tensor(descriptor_input).to(torch.float32)
                pred_scaled = item["model"](descriptor_tensor).detach().numpy()
                model_preds.append(item["train_norm_y"].inverse_transform(pred_scaled).reshape(-1))
        prediction_parts.append(np.vstack(model_preds).mean(axis=0))
    return np.concatenate(prediction_parts)


def conformal_style_empirical_quantile(values, level):
    scores = np.sort(pd.Series(values, dtype=float).dropna().to_numpy())
    if len(scores) == 0:
        raise ValueError("Cannot compute conformal-style empirical quantile from an empty score array.")
    rank = math.ceil((len(scores) + 1) * float(level))
    rank = min(max(rank, 1), len(scores))
    return float(scores[rank - 1]), int(rank), int(len(scores))


def candidate_benchmark(config):
    threshold = float(config["candidate_benchmark_E10_V"])
    if not np.isfinite(threshold):
        raise ValueError("candidate_benchmark_E10_V must be finite.")
    return {
        "threshold_V": threshold,
        "source": "candidate_benchmark_E10_V",
        "high_potential_criterion": "PI95_lower < candidate_benchmark_E10_V",
    }


def load_descriptor_uncertainty_summary(config):
    _, result_dir = run_dirs(config)
    path = result_dir / "descriptor_uncertainty_summary.json"
    if not path.exists():
        raise FileNotFoundError(
            f"Missing descriptor propagation summary: {path}\n"
            "Run 04_propagate_descriptor_uncertainty.py before Step 05."
        )
    with open(path, encoding="utf-8") as f:
        summary = json.load(f)
    if "sigma_descriptor_global_V" not in summary:
        raise ValueError(f"{path} does not contain sigma_descriptor_global_V.")
    return summary


def build_uncertainty_table(config):
    _, result_dir = run_dirs(config)
    grid_ad = pd.read_excel(result_dir / "applicability_domain_grid.xlsx")
    oof = pd.read_excel(result_dir / "cv_seed_oof_predictions.xlsx")
    with open(result_dir / "epistemic_threshold_summary.json", encoding="utf-8") as f:
        epistemic_summary = json.load(f)
    tau_epi = float(epistemic_summary["tau_epi"])
    repeatability_sd, repeatability_data = pooled_repeatability_std(config)
    benchmark = candidate_benchmark(config)
    descriptor_summary = load_descriptor_uncertainty_summary(config)
    descriptor_sigma = float(descriptor_summary["sigma_descriptor_global_V"])

    table = merge_central_predictions(grid_ad, config)
    table["tau_epi"] = tau_epi
    table["ML_epi_class"] = np.where(table["sigma_ML_epi"] <= tau_epi, "low ML epistemic", "high ML epistemic")
    table["AD_flag"] = table["AD95_ilr"].map(ad_flag_from_label)
    table["U_flag"] = np.where(table["sigma_ML_epi"] <= tau_epi, "U-low", "U-high")
    table["Zone"] = [zone_label(ad, unc) for ad, unc in zip(table["AD_flag"], table["U_flag"])]
    table["accepted_by_oof_calibration"] = (
        (table["AD95_ilr"] == AD95_LABEL_INSIDE) & (table["sigma_ML_epi"] <= tau_epi)
    )
    table["sigma_descriptor"] = descriptor_sigma
    oof["sigma_descriptor_oof"] = descriptor_sigma
    oof = prepare_empirical_residual(oof, repeatability_sd)
    table["sigma_emp_residual"] = local_smooth_from_experiments(
        table, oof, config, "sigma_residual_with_repeatability_floor"
    )
    table["sigma_total"] = np.sqrt(
        table["sigma_ML_epi"] ** 2 + table["sigma_emp_residual"] ** 2 + table["sigma_descriptor"] ** 2
    )
    total_variance = table["sigma_total"] ** 2
    table["variance_fraction_ML_epi"] = table["sigma_ML_epi"] ** 2 / total_variance
    table["variance_fraction_descriptor"] = table["sigma_descriptor"] ** 2 / total_variance
    table["variance_fraction_emp_residual"] = table["sigma_emp_residual"] ** 2 / total_variance
    oof["sigma_total_oof"] = np.sqrt(
        oof["sigma_ML_epi_oof"] ** 2
        + oof["sigma_residual_with_repeatability_floor"] ** 2
        + oof["sigma_descriptor_oof"] ** 2
    )
    standardized = oof["abs_residual"] / (oof["sigma_total_oof"] + float(config["eps"]))
    q_value, q_rank, q_n = conformal_style_empirical_quantile(
        standardized, config["standardized_residual_quantile"]
    )
    oof["standardized_abs_residual"] = standardized
    oof.to_excel(result_dir / "cv_seed_oof_predictions_with_residual_proxy.xlsx", index=False)
    component_summary = pd.DataFrame(
        [
            {
                "repeatability_sd_floor_V": repeatability_sd,
                "sigma_descriptor_global_V": descriptor_sigma,
                "OOF_rows": int(len(oof)),
                "residual_smoothing_k": int(config["residual_smoothing_k"]),
            }
        ]
    )
    save_formatted_workbook(
        result_dir / "repeatability_and_residual_audit.xlsx",
        {
            "repeatability": repeatability_data,
            "oof_residual_proxy": oof,
            "component_summary": component_summary,
        },
    )
    table["q0.95"] = q_value
    table["H95"] = q_value * table["sigma_total"]
    table["PI95_lower"] = table["predicted_potential"] - table["H95"]
    table["PI95_upper"] = table["predicted_potential"] + table["H95"]
    table["candidate_benchmark_E10_V"] = benchmark["threshold_V"]
    table.to_excel(result_dir / "uncertainty_table.xlsx", index=False)
    save_json(
        {
            "run_id": config["run_id"],
            "repeatability_sd_floor": repeatability_sd,
            "repeatability_rows": int(len(repeatability_data)),
            "repeatability_source": {
                "path": str(config["repeatability_path"]),
                "sheet": str(config.get("repeatability_sheet", "1")),
                "replicate_columns": ["E10_rep1", "E10_rep2", "E10_rep3"],
            },
            "descriptor_uncertainty": descriptor_summary,
            "empirical_residual_proxy_method": (
                "max(abs_residual^2 - sigma_ML_epi_oof^2 - sigma_descriptor_oof^2, 0), "
                "then apply pooled repeatability SD floor"
            ),
            "q0.95": q_value,
            "q0.95_method": "OOF standardized-residual conformal-style empirical calibration; no strict finite-sample conformal guarantee is claimed because OOF residuals are used instead of an independent calibration set.",
            "q0.95_rank": q_rank,
            "q0.95_score_count": q_n,
            "benchmark": benchmark,
            "epistemic_threshold_summary": epistemic_summary,
            "accepted_by_oof_calibration_counts": {
                str(key): int(value) for key, value in table["accepted_by_oof_calibration"].value_counts().to_dict().items()
            },
            "sigma_ML_epi_quantiles": table["sigma_ML_epi"].quantile([0.25, 0.5, 0.75, 0.95]).to_dict(),
            "sigma_emp_residual_quantiles": table["sigma_emp_residual"].quantile([0.25, 0.5, 0.75, 0.95]).to_dict(),
            "sigma_descriptor_quantiles": table["sigma_descriptor"].quantile([0.25, 0.5, 0.75, 0.95]).to_dict(),
            "sigma_total_quantiles": table["sigma_total"].quantile([0.25, 0.5, 0.75, 0.95]).to_dict(),
            "variance_fraction_ML_epi_quantiles": table["variance_fraction_ML_epi"].quantile(
                [0.25, 0.5, 0.75, 0.95]
            ).to_dict(),
            "variance_fraction_descriptor_quantiles": table["variance_fraction_descriptor"].quantile(
                [0.25, 0.5, 0.75, 0.95]
            ).to_dict(),
            "variance_fraction_emp_residual_quantiles": table[
                "variance_fraction_emp_residual"
            ].quantile([0.25, 0.5, 0.75, 0.95]).to_dict(),
            "variance_fraction_definition": "component variance divided by sigma_total squared; three fractions sum to 1 for every grid point",
            "config": config,
        },
        result_dir / "uncertainty_table_summary.json",
    )
    return result_dir


def ilr_pca_coordinates(table, config, return_details=False):
    experiment_x, experiment_y = load_xy("experiment.xlsx", "metals", "potential")
    experiment_ratios = experiment_x[RATIO_COLUMNS].to_numpy(dtype=float)
    multicomponent_mask = (experiment_ratios > 0).sum(axis=1) > 1
    if int(multicomponent_mask.sum()) != 102:
        raise ValueError(
            "The ILR-PCA visualization reference must contain exactly 102 multicomponent "
            f"experimental compositions; found {int(multicomponent_mask.sum())}."
        )

    delta = float(config["ilr_zero_replacement_delta"])
    experiment_ilr = ilr_transform(experiment_ratios, delta)
    reference_ilr = experiment_ilr[multicomponent_mask]
    grid_ilr = ilr_transform(table[RATIO_COLUMNS].to_numpy(dtype=float), delta)

    pca = PCA(n_components=2, svd_solver="full")
    pca.fit(reference_ilr)

    # Fix the otherwise arbitrary PCA sign by making the largest absolute
    # element-space loading positive for each component.
    basis = ilr_basis(len(RATIO_COLUMNS))
    element_loadings = pca.components_ @ basis
    for component_index in range(pca.n_components_):
        pivot = int(np.argmax(np.abs(element_loadings[component_index])))
        if element_loadings[component_index, pivot] < 0:
            pca.components_[component_index] *= -1.0
            element_loadings[component_index] *= -1.0

    grid_coords = pca.transform(grid_ilr)
    if not return_details:
        return grid_coords

    experiment_coords = pca.transform(experiment_ilr)
    experiment_points = experiment_x[RATIO_COLUMNS].copy().reset_index(drop=True)
    experiment_points.insert(0, "Exp_ID", [f"Exp-{index}" for index in range(1, len(experiment_points) + 1)])
    experiment_points["E10"] = experiment_y.iloc[:, 0].to_numpy(dtype=float)
    experiment_points["point_type"] = np.where(
        multicomponent_mask,
        "multicomponent",
        "single-element",
    )
    experiment_points["ILR_PC1"] = experiment_coords[:, 0]
    experiment_points["ILR_PC2"] = experiment_coords[:, 1]

    metadata_rows = [
        {
            "item": "fit_reference",
            "value": "102 multicomponent experimental compositions",
            "PC1": np.nan,
            "PC2": np.nan,
            "description": "PCA fitting set after ILR transformation.",
        },
        {
            "item": "fit_point_count",
            "value": int(multicomponent_mask.sum()),
            "PC1": np.nan,
            "PC2": np.nan,
            "description": "Number of multicomponent experimental compositions used to fit PCA.",
        },
        {
            "item": "projected_experiment_count",
            "value": int(len(experiment_points)),
            "PC1": np.nan,
            "PC2": np.nan,
            "description": "All experimental points projected with the fixed PCA transformation.",
        },
        {
            "item": "projected_single_element_count",
            "value": int((~multicomponent_mask).sum()),
            "PC1": np.nan,
            "PC2": np.nan,
            "description": "Single-element points excluded from fitting but retained in projection.",
        },
        {
            "item": "ilr_zero_replacement_delta",
            "value": delta,
            "PC1": np.nan,
            "PC2": np.nan,
            "description": "Multiplicative zero-replacement parameter used before ILR transformation.",
        },
        {
            "item": "explained_variance_ratio",
            "value": float(pca.explained_variance_ratio_.sum()),
            "PC1": float(pca.explained_variance_ratio_[0]),
            "PC2": float(pca.explained_variance_ratio_[1]),
            "description": "Individual and cumulative explained-variance ratios of the two displayed PCs.",
        },
    ]
    for ilr_index in range(pca.components_.shape[1]):
        metadata_rows.append(
            {
                "item": f"ILR_coordinate_{ilr_index + 1}_loading",
                "value": np.nan,
                "PC1": float(pca.components_[0, ilr_index]),
                "PC2": float(pca.components_[1, ilr_index]),
                "description": "PCA loading in the four-dimensional ILR coordinate system.",
            }
        )
    for element_index, element in enumerate(RATIO_COLUMNS):
        metadata_rows.append(
            {
                "item": f"element_logcontrast_loading_{element}",
                "value": np.nan,
                "PC1": float(element_loadings[0, element_index]),
                "PC2": float(element_loadings[1, element_index]),
                "description": "Equivalent loading in five-element log-contrast space.",
            }
        )

    details = {
        "experiment_points": experiment_points,
        "metadata": pd.DataFrame(metadata_rows),
        "explained_variance_ratio": pca.explained_variance_ratio_.copy(),
    }
    return grid_coords, details


def plot_reliability_2x4(config):
    import matplotlib.pyplot as plt

    _, result_dir = run_dirs(config)
    table = pd.read_excel(result_dir / "uncertainty_table.xlsx")
    panels = [
        ("predicted_potential", "Predicted potential", "continuous"),
        ("AD95_ilr", "ILR-5NN AD95", "categorical"),
        ("sigma_ML_epi", "ML epistemic uncertainty", "continuous"),
        ("sigma_descriptor", "Descriptor uncertainty", "continuous"),
        ("sigma_emp_residual", "Empirical residual uncertainty", "continuous"),
        ("sigma_total", "Total uncertainty", "continuous"),
        ("H95", "PI95 half-width", "continuous"),
        ("PI95_lower", "Optimistic PI95 bound", "continuous"),
    ]
    zone_export_columns = [
        "AD_flag",
        "tau_epi",
        "U_flag",
        "ML_epi_class",
        "Zone",
        "accepted_by_oof_calibration",
    ]
    require_columns(
        table,
        RATIO_COLUMNS + [panel[0] for panel in panels] + zone_export_columns,
        "uncertainty_table.xlsx",
    )
    coords, pca_details = ilr_pca_coordinates(table, config, return_details=True)

    grid_export_columns = RATIO_COLUMNS + [panel[0] for panel in panels] + zone_export_columns
    grid_export = table[grid_export_columns].copy()
    grid_export.insert(len(RATIO_COLUMNS), "ILR_PC1", coords[:, 0])
    grid_export.insert(len(RATIO_COLUMNS) + 1, "ILR_PC2", coords[:, 1])
    zone_definitions = pd.DataFrame(
        [
            {
                "Zone": "Z1",
                "AD_flag": "AD-in",
                "U_flag": "U-low",
                "definition": "Inside ILR-5NN AD95 and sigma_ML_epi <= tau_epi",
            },
            {
                "Zone": "Z2",
                "AD_flag": "AD-out",
                "U_flag": "U-low",
                "definition": "Outside ILR-5NN AD95 and sigma_ML_epi <= tau_epi",
            },
            {
                "Zone": "Z3",
                "AD_flag": "AD-in",
                "U_flag": "U-high",
                "definition": "Inside ILR-5NN AD95 and sigma_ML_epi > tau_epi",
            },
            {
                "Zone": "Z4",
                "AD_flag": "AD-out",
                "U_flag": "U-high",
                "definition": "Outside ILR-5NN AD95 and sigma_ML_epi > tau_epi",
            },
        ]
    )
    with pd.ExcelWriter(result_dir / "reliability_2x4_plot_data.xlsx", engine="openpyxl") as writer:
        grid_export.to_excel(writer, sheet_name="grid", index=False)
        pca_details["experiment_points"].to_excel(writer, sheet_name="experiment_points", index=False)
        pca_details["metadata"].to_excel(writer, sheet_name="pca_metadata", index=False)
        zone_definitions.to_excel(writer, sheet_name="zone_definitions", index=False)

    fig, axes = plt.subplots(2, 4, figsize=(22, 10), constrained_layout=True)
    fig.suptitle("ILR-PCA fitted on 102 multicomponent experimental compositions", fontsize=13)
    for ax, (column, title, kind) in zip(axes.reshape(-1), panels):
        if kind == "continuous":
            scatter = ax.scatter(coords[:, 0], coords[:, 1], c=table[column], s=2, cmap="viridis")
            fig.colorbar(scatter, ax=ax, label=column)
        else:
            categories = sorted(table[column].dropna().unique())
            colors = plt.cm.Set1(np.linspace(0, 1, max(len(categories), 3)))
            for idx, category in enumerate(categories):
                mask = table[column] == category
                ax.scatter(coords[mask, 0], coords[mask, 1], s=2, color=colors[idx], label=str(category), alpha=0.8)
            ax.legend(markerscale=4, fontsize=8)
        ax.set_title(title)
        ax.set_xlabel("ILR-PC1 (102-experiment reference)")
        ax.set_ylabel("ILR-PC2 (102-experiment reference)")
    fig.savefig(
        result_dir / "reliability_2x4_map.png",
        dpi=300,
        bbox_inches="tight",
        metadata={"Software": None},
    )
    plt.close(fig)
    return result_dir


def composition_key_frame(frame):
    return frame[RATIO_COLUMNS].round(8).astype(str).agg("|".join, axis=1)


def load_candidate_table(config):
    _, result_dir = run_dirs(config)
    formal_path = result_dir / "uncertainty_table.xlsx"
    if not formal_path.exists():
        raise FileNotFoundError(
            f"Missing formal uncertainty table: {formal_path}\n"
            "Run Steps 01-05 before candidate recommendation. Temporary or estimated prediction intervals are not used."
        )
    table = pd.read_excel(formal_path)
    required = [
        "predicted_potential",
        "sigma_ML_epi",
        "sigma_descriptor",
        "sigma_emp_residual",
        "sigma_total",
        "H95",
        "PI95_lower",
        "PI95_upper",
        "Zone",
    ]
    missing = [column for column in required if column not in table.columns]
    if missing:
        raise ValueError(
            f"{formal_path} is missing formal uncertainty columns: {missing}. "
            "Run Steps 01-05 before candidate recommendation."
        )
    numeric = [column for column in required if column != "Zone"]
    if table[numeric].isna().any().any():
        raise ValueError(f"{formal_path} contains missing formal uncertainty values.")
    if not np.allclose(table["PI95_lower"], table["predicted_potential"] - table["H95"], atol=1e-12):
        raise ValueError("PI95_lower does not equal predicted_potential - H95 in uncertainty_table.xlsx.")
    if not np.allclose(table["PI95_upper"], table["predicted_potential"] + table["H95"], atol=1e-12):
        raise ValueError("PI95_upper does not equal predicted_potential + H95 in uncertainty_table.xlsx.")
    return table


def attach_candidate_benchmark(table, config):
    benchmark = candidate_benchmark(config)
    out = table.copy()
    out["candidate_benchmark_E10_V"] = benchmark["threshold_V"]
    return out, benchmark


def experiment_exclusion_keys():
    experiment_x, _ = load_xy("experiment.xlsx", "metals", "potential")
    return set(composition_key_frame(experiment_x))


def ordinary_experiment_points():
    experiment_x, _ = load_xy("experiment.xlsx", "metals", "potential")
    nonzero_count = (experiment_x[RATIO_COLUMNS] > 0).sum(axis=1)
    return experiment_x.loc[nonzero_count > 1, RATIO_COLUMNS].copy()


def select_global_prediction_top(table, config, excluded_keys):
    count = int(config.get("global_top_count", 10))
    if count <= 0:
        raise ValueError("global_top_count must be a positive integer.")
    pool = filter_excluded(table, excluded_keys)
    selected = pool.sort_values(["predicted_potential"] + RATIO_COLUMNS, kind="mergesort").head(count).copy()
    if len(selected) != count:
        raise ValueError(f"Global prediction selection returned {len(selected)} rows; expected {count}.")
    selected["global_prediction_rank"] = np.arange(1, count + 1)
    return selected


def filter_excluded(table, keys):
    out = table.copy()
    out["_key"] = composition_key_frame(out)
    return out.loc[~out["_key"].isin(keys)].drop(columns=["_key"]).copy()


def _reference_distances(coords, selected_indices, existing_coords, remaining):
    reference = coords[selected_indices]
    if existing_coords is not None:
        reference = np.vstack([reference, existing_coords]) if len(reference) else existing_coords
    return euclidean_distance_matrix(coords[remaining], reference).min(axis=1)


def ilr_diverse_select_high_potential(pool, count, config, existing=None):
    if count <= 0 or pool.empty:
        return pool.head(0).copy()
    work = pool.copy().reset_index(drop=True)
    coords = ilr_transform(work[RATIO_COLUMNS].to_numpy(float), config["ilr_zero_replacement_delta"])
    selected_indices = []
    existing_coords = None
    if existing is not None and len(existing):
        existing_coords = ilr_transform(existing[RATIO_COLUMNS].to_numpy(float), config["ilr_zero_replacement_delta"])
    optimistic = work["PI95_lower"].to_numpy(dtype=float)
    work["ilr_distance_score"] = np.nan
    work["optimistic_PI_score"] = np.nan
    work["selection_score"] = np.nan
    for _ in range(min(count, len(work))):
        remaining = [idx for idx in range(len(work)) if idx not in selected_indices]
        if not selected_indices and existing_coords is None:
            choice = min(remaining, key=lambda idx: optimistic[idx])
            work.loc[choice, "optimistic_PI_score"] = 1.0
            work.loc[choice, "selection_score"] = 0.45
        else:
            distances = _reference_distances(coords, selected_indices, existing_coords, remaining)
            dist_score = distances / (distances.max() + 1e-12)
            pred_score = 1.0 - (optimistic[remaining] - optimistic[remaining].min()) / (
                optimistic[remaining].max() - optimistic[remaining].min() + 1e-12
            )
            score = 0.55 * dist_score + 0.45 * pred_score
            position = int(np.argmax(score))
            choice = remaining[position]
            work.loc[choice, "ilr_distance_score"] = float(dist_score[position])
            work.loc[choice, "optimistic_PI_score"] = float(pred_score[position])
            work.loc[choice, "selection_score"] = float(score[position])
        selected_indices.append(choice)
    return work.iloc[selected_indices].copy()


def format_candidate_output(frame):
    out = frame.copy()
    preferred = [
        "Fe",
        "Co",
        "Ni",
        "Mn",
        "Zn",
        "predicted_potential",
        "sigma_ML_epi",
        "sigma_descriptor",
        "sigma_emp_residual",
        "sigma_total",
        "H95",
        "PI95_lower",
        "PI95_upper",
        "Reason",
        "Region",
        "global_prediction_rank",
        "zone_rank",
        "global_rank",
        "output_order",
        "selection_stage",
        "candidate_benchmark_E10_V",
        "candidate_eligible",
        "formal_PI95_eligible",
        "ilr_distance_score",
        "optimistic_PI_score",
        "selection_score",
        "distance_reference",
        "uncertainty_source",
    ]
    extras = [col for col in out.columns if col not in preferred]
    return out[[col for col in preferred if col in out.columns] + extras]


def add_global_pi_rank(frame, full_table):
    ranked = full_table[RATIO_COLUMNS + ["PI95_lower"]].copy()
    ranked["_key"] = composition_key_frame(ranked)
    ranked = ranked.sort_values(["PI95_lower"] + RATIO_COLUMNS).reset_index(drop=True)
    ranked["global_rank"] = np.arange(1, len(ranked) + 1)
    rank_map = ranked.drop_duplicates("_key").set_index("_key")["global_rank"]
    out = frame.copy()
    out["_key"] = composition_key_frame(out)
    out["global_rank"] = out["_key"].map(rank_map).astype("Int64")
    return out.drop(columns=["_key"])


def add_zone_pi_rank(frame, full_table):
    ranked = full_table[RATIO_COLUMNS + ["Zone", "PI95_lower"]].copy()
    ranked["_key"] = composition_key_frame(ranked)
    ranked = ranked.sort_values(["Zone", "PI95_lower"] + RATIO_COLUMNS).reset_index(drop=True)
    ranked["zone_rank"] = ranked.groupby("Zone").cumcount() + 1
    rank_map = ranked.drop_duplicates("_key").set_index("_key")["zone_rank"]
    out = frame.copy()
    out["_key"] = composition_key_frame(out)
    out["zone_rank"] = out["_key"].map(rank_map).astype("Int64")
    return out.drop(columns=["_key"])


def order_high_potential_output(frame):
    if frame.empty:
        return frame.copy()
    zone_order = {"Z1": 1, "Z2": 2, "Z3": 3, "Z4": 4}
    stage_order = {"global_top10": 0, "zone_top": 1, "zone_diversity": 2}
    out = frame.copy()
    out["_zone_order"] = out["Region"].map(zone_order).fillna(99)
    out["_stage_order"] = out["selection_stage"].map(stage_order).fillna(99)
    out = out.sort_values(
        ["_stage_order", "global_prediction_rank", "_zone_order", "zone_rank", "global_rank"],
        na_position="last",
    ).reset_index(drop=True)
    out["output_order"] = np.arange(1, len(out) + 1)
    return out.drop(columns=["_zone_order", "_stage_order"])

def recommend_high_potential_candidates(config):
    _, result_dir = run_dirs(config)
    table = load_candidate_table(config)
    table, benchmark = attach_candidate_benchmark(table, config)
    rank_reference = table.copy()
    threshold = benchmark["threshold_V"]
    experiment_keys = experiment_exclusion_keys()

    global_top = select_global_prediction_top(table, config, experiment_keys)
    global_top["Reason"] = "Global predicted-potential Top-10"
    global_top["Region"] = global_top["Zone"]
    global_top["selection_stage"] = "global_top10"
    global_top["distance_reference"] = "not used for global Top-10"
    global_top["target_quota"] = int(config.get("global_top_count", 10))
    global_top["eligible_zone_count"] = np.nan
    global_top["formal_PI95_eligible"] = global_top["PI95_lower"] < threshold
    global_top["candidate_eligible"] = True

    uncertainty_exclusions = set(experiment_keys)
    uncertainty_exclusions.update(composition_key_frame(global_top))
    uncertainty_pool = filter_excluded(table, uncertainty_exclusions)
    uncertainty_pool = add_zone_pi_rank(uncertainty_pool, rank_reference)
    eligible = uncertainty_pool.loc[uncertainty_pool["PI95_lower"] < threshold].copy()
    top_quota = {"Z2": 3, "Z1": 2, "Z3": 2}
    diversity_quota = {"Z4": 8, "Z2": 7, "Z1": 4, "Z3": 4}
    total_quota = {
        zone: top_quota.get(zone, 0) + diversity_quota.get(zone, 0)
        for zone in ["Z1", "Z2", "Z3", "Z4"]
    }
    shortages = {
        zone: {"required": count, "eligible": int((eligible["Zone"] == zone).sum())}
        for zone, count in total_quota.items()
        if int((eligible["Zone"] == zone).sum()) < count
    }
    if shortages:
        raise ValueError(
            "Insufficient candidates satisfying PI95_lower < "
            f"{threshold:.4f} V after exclusions: {shortages}. "
            "Candidates outside the formal prediction-interval criterion will not be used as fallback."
        )
    ordinary_exp = ordinary_experiment_points()
    selected_parts = []
    top_by_zone = {}
    zone_diagnostics = []

    for zone in ["Z2", "Z1", "Z3"]:
        pool = eligible.loc[eligible["Zone"] == zone].sort_values(
            ["PI95_lower"] + RATIO_COLUMNS
        ).copy()
        top = pool.head(top_quota[zone]).copy()
        top["Reason"] = f"{zone} formal PI95 optimistic-bound top"
        top["Region"] = zone
        top["selection_stage"] = "zone_top"
        top["distance_reference"] = "not used for zone top"
        top["target_quota"] = top_quota[zone]
        top["eligible_zone_count"] = int((eligible["Zone"] == zone).sum())
        selected_parts.append(top)
        top_by_zone[zone] = top

    for zone in ["Z4", "Z2", "Z1", "Z3"]:
        pool = eligible.loc[eligible["Zone"] == zone].sort_values(
            ["PI95_lower"] + RATIO_COLUMNS
        ).copy()
        if zone in top_by_zone and not top_by_zone[zone].empty:
            pool = filter_excluded(pool, set(composition_key_frame(top_by_zone[zone])))
        reference = pd.concat(
            [
                ordinary_exp,
                global_top[RATIO_COLUMNS],
                top_by_zone.get(zone, uncertainty_pool.head(0))[RATIO_COLUMNS],
            ],
            ignore_index=True,
        )
        distance_reference = f"{len(ordinary_exp)} experiments + global Top-10 + {zone} top points"
        diverse = ilr_diverse_select_high_potential(pool, diversity_quota[zone], config, existing=reference)
        if len(diverse) != diversity_quota[zone]:
            raise ValueError(
                f"{zone} diversity selection returned {len(diverse)} rows; expected {diversity_quota[zone]}."
            )
        diverse["Reason"] = f"{zone} formal PI95 eligibility + ILR diversity"
        diverse["Region"] = zone
        diverse["selection_stage"] = "zone_diversity"
        diverse["distance_reference"] = distance_reference
        diverse["target_quota"] = diversity_quota[zone]
        diverse["eligible_zone_count"] = int((eligible["Zone"] == zone).sum())
        selected_parts.append(diverse)
        zone_diag = {
            "zone": zone,
            "zone_total_count": int((uncertainty_pool["Zone"] == zone).sum()),
            "formal_PI_eligible_count": int((eligible["Zone"] == zone).sum()),
            "top_selected_count": int(len(top_by_zone.get(zone, []))),
            "diversity_selected_count": int(len(diverse)),
        }
        zone_diagnostics.append(zone_diag)
        print(
            f"{zone}: zone_total={zone_diag['zone_total_count']}, "
            f"formal_PI_eligible={zone_diag['formal_PI_eligible_count']}, "
            f"top_selected={zone_diag['top_selected_count']}, "
            f"diversity_selected={zone_diag['diversity_selected_count']}"
        )
    selected = pd.concat(selected_parts, ignore_index=True)
    expected_zone_counts = {"Z1": 6, "Z2": 10, "Z3": 6, "Z4": 8}
    actual_zone_counts = {str(key): int(value) for key, value in selected["Region"].value_counts().to_dict().items()}
    if len(selected) != 30 or actual_zone_counts != expected_zone_counts:
        raise ValueError(
            f"High-potential selection quota mismatch: rows={len(selected)}, zones={actual_zone_counts}; "
            f"expected rows=30, zones={expected_zone_counts}."
        )
    if selected.duplicated(RATIO_COLUMNS).any():
        raise ValueError("Uncertainty-guided selection contains duplicate compositions.")
    if not (selected["PI95_lower"] < threshold).all():
        raise ValueError("All 30 uncertainty-guided candidates must satisfy the formal PI95 criterion.")

    selected = add_global_pi_rank(selected, rank_reference)
    selected = add_zone_pi_rank(selected, rank_reference)
    selected["formal_PI95_eligible"] = True
    selected["candidate_eligible"] = True

    global_top = add_global_pi_rank(global_top, rank_reference)
    global_top = add_zone_pi_rank(global_top, rank_reference)
    combined = pd.concat([global_top, selected], ignore_index=True, sort=False)
    if len(combined) != 40 or combined.duplicated(RATIO_COLUMNS).any():
        raise ValueError("Combined candidate selection must contain 40 unique compositions.")
    if set(composition_key_frame(combined)) & experiment_keys:
        raise ValueError("Combined candidate selection overlaps the existing experimental data.")
    combined = order_high_potential_output(combined)
    combined["uncertainty_source"] = "result/uncertainty_table.xlsx"
    output = format_candidate_output(combined)
    output.to_excel(result_dir / "high_potential_candidate_recommendations.xlsx", index=False)
    save_json(
        {
            "run_id": config["run_id"],
            "uncertainty_source": "result/uncertainty_table.xlsx",
            "benchmark": benchmark,
            "rows": int(len(output)),
            "global_prediction_top_count": int((output["selection_stage"] == "global_top10").sum()),
            "uncertainty_guided_count": int((output["selection_stage"] != "global_top10").sum()),
            "zone_counts": {str(k): int(v) for k, v in output["Region"].value_counts().to_dict().items()},
            "selection_stage_counts": {
                str(k): int(v) for k, v in output["selection_stage"].value_counts().to_dict().items()
            },
            "zone_diagnostics": zone_diagnostics,
            "selection_channels": {
                "global_top10": "10 lowest predicted_potential values among untested grid points; no PI95 gate",
                "uncertainty_guided_30": "PI95_lower < candidate_benchmark_E10_V, followed by zone quotas and ILR diversity",
            },
            "global_rank_definition": "rank of PI95_lower among all grid points",
            "zone_rank_definition": "rank of PI95_lower within the same Zone",
            "global_prediction_rank_definition": "rank of predicted_potential among untested grid points",
            "scoring_formula": "score = 0.55 * ILR_distance_score + 0.45 * optimistic_PI95_lower_score",
        },
        result_dir / "high_potential_candidate_recommendations_summary.json",
    )
    return result_dir
